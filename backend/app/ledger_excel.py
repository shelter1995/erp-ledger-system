from __future__ import annotations

from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from .department_service import department_filter
from sqlalchemy import text, bindparam
from sqlalchemy.engine import Connection

from .edit_versions import line_context
from .history_queries import order_match, manager_match
from .auth import CurrentUser, apply_department_scope
from .serializers import clean_value


TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "市场部业务台账模板.xlsx"
TEMPLATE_FILE_NAME = "市场部业务台账模板.xlsx"
EXPORT_FILE_NAME = "市场部业务台账.xlsx"
LEGACY_TEMPLATE_HEADERS = [
    "全额/净额", "项目编号", "部门", "分公司", "客户经理", "订单日期", "业务类型", "统计类别",
    "三级团队名称", "客户单位名称", "用户", "区域平台", "订单号", "项目名称", "货物名称", "规格型号",
    "单位", "数量", "销售税率", "不含税单价", "单价", "不含税收入", "订单价值", "采购厂商", "采购税率",
    "不含税采购单价", "采购单价", "不含税成本", "采购金额", "交付日期", "交付数量", "交付不含税收入", "交付价值",
    "交付不含税成本", "交付成本", "待交付数量", "待交付金额（不含税）", "待交付金额", "公司合同号",
    "付款期限", "履行期限", "合同签订金额", "待签合同金额", "收票日期", "发票号码", "收票金额",
    "入库日期", "凭证号", "入库金额", "入账日期", "凭证号", "入账金额", "待入账金额", "到期付款日",
    "付款日期", "付款凭证号", "付款金额", "付款日期", "付款凭证号", "付款金额", "付款金额", "应付帐款",
    "不含税毛利润", "税金", "退税", "毛利润", "毛利率", "合同签订日期", "公司合同号", "合同价值",
    "履行期限", "待签合同金额", "开票单据号", "开票日期", "发票号", "发票金额", "待开发票金额",
    "已交付未开票", "回款日期", "缴款单号", "回款金额", "回款占比", "回款日期", "缴款单号",
    "回款金额", "回款占比", "回款合计", "应收款", "是否关闭", "人工成本", "其他成本",
]
TEMPLATE_HEADERS = list(LEGACY_TEMPLATE_HEADERS)
RENAMED_HEADERS = {
    6: "销售订单日期", 13: "销售订单号", 15: "物资/服务名称",
    22: "不含税订单金额", 23: "销售订单金额", 28: "不含税采购金额",
    29: "采购金额", 33: "交付收入", 69: "翔云合同号", 70: "合同金额",
}
for _column, _label in RENAMED_HEADERS.items():
    TEMPLATE_HEADERS[_column - 1] = _label
TEMPLATE_HEADERS[87:88] = ["交付应收款", "开票应收款"]


def standard_template_version(headers) -> int | None:
    names = [str(value or "").strip() for value in headers]
    while names and not names[-1]:
        names.pop()
    if names == TEMPLATE_HEADERS:
        return 92
    if names == LEGACY_TEMPLATE_HEADERS:
        return 91
    return None


def normalize_standard_row(values, version: int):
    """Only a verified header layout may shift financial columns."""
    row = list(values[:version]) + [None] * max(0, version - len(values))
    if version == 91:
        row.insert(88, None)
    return row


def _load_current_template():
    """Upgrade the bundled layout in memory; never import reference sample data."""
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook["Sheet1"]
    version = standard_template_version([cell.value for cell in worksheet[2]])
    if version not in (91, 92):
        workbook.close()
        raise ValueError("内置台账模板表头不受支持")
    if worksheet.max_row > 2:
        worksheet.delete_rows(3, worksheet.max_row - 2)
    if version == 91:
        worksheet.unmerge_cells("CI1:CK1")
        worksheet.unmerge_cells("CL1:CM1")
        worksheet.insert_cols(89)
        for row_no in (1, 2):
            worksheet.cell(row_no, 89)._style = copy(worksheet.cell(row_no, 88)._style)
        worksheet.column_dimensions["CN"].width = worksheet.column_dimensions["CM"].width
        worksheet.merge_cells("CI1:CL1")
        worksheet.merge_cells("CM1:CN1")
    for column, label in enumerate(TEMPLATE_HEADERS, 1):
        worksheet.cell(2, column).value = label
        if column in RENAMED_HEADERS:
            worksheet.cell(2, column).fill = PatternFill(fill_type="solid", fgColor="FFFEFF00")
    for coordinate, label in {
        "AM1": "采购合同（杨志毅）", "AR1": "收票情况（杨志毅）",
        "BB1": "一期付款（杨志毅）", "BF1": "二期付款（杨志毅）",
        "CM1": "补充成本",
    }.items():
        worksheet[coordinate] = label
    worksheet.print_area = "A1:CN3"
    return workbook
EDITABLE_ORDER_KEYS = [
    "amount_type", "project_code", "department", "branch_company", "account_manager", "order_date",
    "business_type", "statistical_category", "team_name", "customer_unit_name", "user_name",
    "regional_platform", "order_no", "project_name", "goods_name", "specification_model", "unit_name",
    "quantity", "sales_tax_rate", "net_unit_price", "unit_price", "net_revenue", "order_value",
]
PURCHASE_EDITOR_KEYS_BY_COLUMN = {
    24: "supplier_name",
    25: "purchase_tax_rate",
    26: "purchase_unit_price_no_tax",
    27: "purchase_unit_price",
    28: "cost_no_tax",
    29: "purchase_amount",
    30: "delivery_date",
    31: "delivery_quantity",
    32: "delivery_revenue_no_tax",
    33: "delivery_value",
    34: "delivery_cost_no_tax",
    35: "delivery_cost",
    36: "pending_delivery_quantity",
    37: "pending_delivery_amount_no_tax",
    38: "pending_delivery_amount",
    39: "purchase_contract_no",
    40: "payment_terms",
    41: "purchase_performance_period",
    42: "purchase_signed_amount",
    43: "purchase_unsigned_amount",
    44: "received_invoice_date",
    45: "purchase_invoice_no",
    46: "purchase_invoice_amount",
    47: "warehouse_date",
    48: "warehouse_voucher_no",
    49: "warehouse_amount",
    50: "booked_date",
    51: "booked_voucher_code",
    52: "booked_amount",
    53: "pending_booked_amount",
    54: "payment1_due_date",
    55: "payment1_date",
    56: "payment1_voucher_no",
    57: "payment1_amount",
    58: "payment2_date",
    59: "payment2_voucher_no",
    60: "payment2_amount",
    61: "total_paid",
    62: "accounts_payable",
    63: "gross_profit_no_tax",
    64: "tax_difference",
    65: "tax_refund",
    66: "gross_profit",
    67: "gross_profit_margin_no_tax",
}
PURCHASE_EDITABLE_COLUMN_NUMBERS = frozenset(
    (set(range(24, 53)) | set(range(54, 61))) - {28, 29, 32, 33, 34, 35, 36, 37, 38, 43}
)
SALES_EDITOR_KEYS_BY_COLUMN = {
    68: "contract_signed_date",
    69: "sales_contract_no",
    70: "sales_contract_value",
    71: "sales_performance_period",
    72: "sales_unsigned_contract_amount",
    73: "invoice_doc_no",
    74: "invoice_date",
    75: "sales_invoice_no",
    76: "sales_invoice_amount",
    77: "pending_invoice_amount",
    78: "delivered_not_invoiced_amount",
    79: "receipt1_date",
    80: "receipt1_notice_no",
    81: "receipt1_amount",
    82: "receipt1_ratio",
    83: "receipt2_date",
    84: "receipt2_notice_no",
    85: "receipt2_amount",
    86: "receipt2_ratio",
    87: "total_received",
    88: "delivery_accounts_receivable",
    89: "invoice_accounts_receivable",
    90: "close_status",
    91: "labor_cost",
    92: "other_cost",
}
SALES_EDITABLE_COLUMN_NUMBERS = frozenset(
    (set(range(68, 87)) | {91, 92}) - {77, 78, 82, 86}
)
REQUIRED_ORDER_COLUMN_NUMBERS = {2, 10, 13, 15}
SAMPLE_PROJECT_CODE = "示例项目编号-请替换"
SAMPLE_ORDER_NO = "示例销售订单号-请替换"
SAMPLE_TEMPLATE_ROW = [
    "全额", SAMPLE_PROJECT_CODE, "科贸部", "安徽分公司", "示例客户经理", date(2026, 7, 22), "商品销售", "常规业务",
    "示例三级团队", "示例客户单位", "示例最终用户", "示例区域平台", SAMPLE_ORDER_NO, "示例项目名称",
    "示例设备", "EXAMPLE-001", "台", 10, 0.13, 100, 113, 1000, 1130, "示例采购厂商", 0.13, 70, 79.1, 700, 791,
    date(2026, 7, 23), 10, 1000, 1130, 700, 791, 0, 0, 0, "CGHT-EXAMPLE-001", "验收后30日内付款",
    "合同签订后30日内", 791, 0, date(2026, 7, 24), "CGFP-EXAMPLE-001", 791, date(2026, 7, 25),
    "RK-EXAMPLE-001", 791, date(2026, 7, 26), "RZ-EXAMPLE-001", 791, 0, date(2026, 8, 25),
    date(2026, 7, 27), "FK-EXAMPLE-001", 791, None, None, None, 791, 0, 300, 39, 0, 339, 0.3,
    date(2026, 7, 22), "XSHT-EXAMPLE-001", 1130, "合同签订后30日内", 0, "KP-EXAMPLE-001",
    date(2026, 7, 28), "FP-EXAMPLE-001", 1130, 0, 0, date(2026, 7, 29), "JK-EXAMPLE-001", 1130, 1,
    None, None, None, None, 1130, 0, 0, "关闭", None, None,
]

if len(SAMPLE_TEMPLATE_ROW) != len(TEMPLATE_HEADERS):
    raise RuntimeError("业务台账示例行字段数量与模板表头不一致")


def content_disposition(file_name: str) -> str:
    return f"attachment; filename*=UTF-8''{quote(file_name)}"


def template_bytes() -> bytes:
    if not TEMPLATE_PATH.is_file():
        raise FileNotFoundError(f"导入模板不存在：{TEMPLATE_PATH}")
    workbook = _load_current_template()
    worksheet = workbook["Sheet1"]
    if worksheet.max_row > 3:
        worksheet.delete_rows(4, worksheet.max_row - 3)
    example_fill = PatternFill(fill_type="solid", fgColor="FFF7D6")
    example_font = Font(color="7C5C00")
    for column_no, value in enumerate(SAMPLE_TEMPLATE_ROW, start=1):
        cell = worksheet.cell(3, column_no, value)
        _copy_header_alignment(worksheet.cell(2, column_no), cell)
        cell.fill = example_fill
        cell.font = example_font
        if column_no in DATE_COLUMNS and isinstance(value, (date, datetime)):
            cell.number_format = "yyyy-mm-dd"
        elif column_no in PERCENT_COLUMNS and value is not None:
            cell.number_format = "0.00%"
        elif column_no in NUMBER_COLUMNS and value is not None:
            cell.number_format = "#,##0.00"
    worksheet.row_dimensions[3].height = 24
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def is_template_sample_row(row: tuple[Any, ...]) -> bool:
    if len(row) < 13:
        return False
    return str(row[1] or "").strip() == SAMPLE_PROJECT_CODE and str(row[12] or "").strip() == SAMPLE_ORDER_NO


def _literal_export_cell(worksheet, row_no: int, column_no: int, value: Any):
    cell = worksheet.cell(row_no, column_no, value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def export_ledger_bytes(
    conn: Connection,
    user: CurrentUser,
    filters: Mapping[str, object] | None = None,
) -> bytes:
    workbook = _load_current_template()
    worksheet = workbook["Sheet1"]
    if worksheet.max_row > 2:
        worksheet.delete_rows(3, worksheet.max_row - 2)

    rows = [dict(row) for row in filtered_export_rows(conn, user, filters)]
    # A multi-delivery line still occupies one Excel row. Export all dates and
    # their positional values rather than silently keeping the last date.
    from .source_delivery import DELIVERY_COLUMNS
    by_line = {}
    ids = [row['order_line_id'] for row in rows]
    from .history_queries import source_order_paths
    paths = source_order_paths(conn, ids)
    for start in range(0, len(ids), 1000):
        phases = conn.execute(text('SELECT * FROM delivery_record WHERE order_line_id IN :ids AND deleted_at IS NULL ORDER BY id')
            .bindparams(bindparam('ids', expanding=True)), {'ids':ids[start:start+1000]}).mappings()
        for phase in phases:
            by_line.setdefault(phase['order_line_id'], []).append(phase)
    for row in rows:
        path = paths.get(row['order_line_id'])
        if path and path[-1] == row['order_no']:
            row['order_no'] = '/'.join(path)
        phases = by_line.get(row['order_line_id'], [])
        if len(phases) > 1:
            row['delivery_date'] = ';'.join(str(p['delivery_date'] or '') for p in phases)
            for col, field in DELIVERY_COLUMNS.items():
                if col < 36 and any(p[field] is not None for p in phases):
                    row[field] = '/'.join(str(p[field]) if p[field] is not None else '' for p in phases)

    for row_no, row in enumerate(rows, start=3):
        values = _row_values(dict(row))
        for column_no, value in enumerate(values, start=1):
            cell = _literal_export_cell(worksheet, row_no, column_no, value)
            _copy_header_alignment(worksheet.cell(2, column_no), cell)
            if column_no in DATE_COLUMNS and isinstance(value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            elif column_no in PERCENT_COLUMNS and value is not None:
                cell.number_format = "0.00%"
            elif column_no in NUMBER_COLUMNS and value is not None:
                cell.number_format = "#,##0.00"

    worksheet.print_area = f"A1:CN{max(3, worksheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def editor_columns(scope: str = "basic") -> list[dict[str, Any]]:
    if scope not in {"basic", "purchase", "sales"}:
        raise ValueError(f"Unsupported editor scope: {scope}")
    editable_columns = (
        PURCHASE_EDITABLE_COLUMN_NUMBERS
        if scope == "purchase"
        else SALES_EDITABLE_COLUMN_NUMBERS
        if scope == "sales"
        else frozenset(range(1, len(EDITABLE_ORDER_KEYS) + 1))
    )
    columns: list[dict[str, Any]] = []
    for column_no, label in enumerate(TEMPLATE_HEADERS, start=1):
        key = (
            EDITABLE_ORDER_KEYS[column_no - 1]
            if column_no <= len(EDITABLE_ORDER_KEYS)
            else (
                SALES_EDITOR_KEYS_BY_COLUMN.get(column_no)
                if scope == "sales"
                else PURCHASE_EDITOR_KEYS_BY_COLUMN.get(column_no)
            )
        )
        columns.append(
            {
                "excel_column": _excel_column_name(column_no),
                "label": label,
                "key": key,
                "value_type": (
                    "date"
                    if column_no in DATE_COLUMNS
                    else "percentage"
                    if column_no in PERCENT_COLUMNS
                    else "number"
                    if column_no in NUMBER_COLUMNS
                    else "text"
                ),
                "editable": column_no in editable_columns and column_no not in {22, 23, 28, 29, 32, 33, 34, 35, 36, 37, 38, 43, 77, 78, 82, 86, 89},
                "required": scope == "basic" and column_no in REQUIRED_ORDER_COLUMN_NUMBERS,
            }
        )
    return columns


def editor_rows_for_order_lines(
    conn: Connection,
    user: CurrentUser,
    order_line_ids: list[int],
) -> list[dict[str, Any]]:
    unique_ids = list(dict.fromkeys(order_line_ids))
    if not unique_ids:
        return []
    params: dict[str, object] = {}
    placeholders: list[str] = []
    for index, order_line_id in enumerate(unique_ids):
        key = f"order_line_id_{index}"
        params[key] = order_line_id
        placeholders.append(f":{key}")
    conditions = [
        "p.deleted_at IS NULL",
        "so.deleted_at IS NULL",
        "ol.deleted_at IS NULL",
        f"ol.id IN ({', '.join(placeholders)})",
    ]
    apply_department_scope(conditions, params, user, "fin.department")
    rows = conn.execute(text(_export_sql(" AND ".join(conditions))), params).mappings().all()
    edit_context = line_context(conn, [r["order_line_id"] for r in rows])
    result: list[dict[str, Any]] = []
    for row in rows:
        mapped = dict(row)
        values = _row_values(mapped)
        # Editors expose the stored two-decimal amounts. Raw source precision
        # is reserved for profit calculation and downloadable source data.
        for index in (22, 28):
            if values[index] is not None:
                values[index] = Decimal(values[index]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        # The editor uses the same percentage unit as the API (13 means 13%),
        # while the downloadable workbook stores 0.13 with percentage formatting.
        for index, key in (
            (18, "sales_tax_rate"),
            (24, "purchase_tax_rate"),
            (66, "gross_profit_margin_no_tax"),
            (81, "receipt1_ratio"),
            (85, "receipt2_ratio"),
        ):
            values[index] = mapped.get(key)
        result.append(
            {
                "order_line_id": int(mapped["order_line_id"]),
                "values": [clean_value(value) for value in values],
                "edit_context": edit_context,
            }
        )
    return result


def editor_payload_snapshot(
    conn: Connection,
    user: CurrentUser,
    order_line_id: int,
    scope: str,
) -> dict[str, Any]:
    rows = editor_rows_for_order_lines(conn, user, [order_line_id])
    if not rows:
        raise LookupError(f"Order line {order_line_id} is unavailable")
    values = rows[0]["values"]
    snapshot: dict[str, Any] = {"order_line_id": order_line_id}
    for index, column in enumerate(editor_columns(scope)):
        if column["editable"] and column["key"]:
            snapshot[str(column["key"])] = values[index]
    return snapshot


def editor_changed_keys(
    scope: str,
    current: dict[str, Any],
    submitted: dict[str, object],
) -> list[str]:
    changed: list[str] = []
    for column in editor_columns(scope):
        key = column["key"]
        if not column["editable"] or not key:
            continue
        if _editor_comparison_value(column["value_type"], current.get(key)) != _editor_comparison_value(
            column["value_type"],
            submitted.get(key),
        ):
            changed.append(str(key))
    return changed


def _editor_comparison_value(value_type: str, value: Any) -> Any:
    if value is None or str(value).strip() == "":
        return None
    if value_type in {"number", "percentage"}:
        try:
            return Decimal(str(value).replace(",", "").strip())
        except InvalidOperation:
            return str(value).strip()
    if value_type == "date":
        if isinstance(value, (date, datetime)):
            return value.isoformat()[:10]
        return str(value).strip().replace("/", "-")
    return str(value).strip()


DATE_COLUMNS = {6, 30, 44, 47, 50, 54, 55, 58, 68, 74, 79, 83}
PERCENT_COLUMNS = {19, 25, 67, 82, 86}
TEXT_COLUMNS_WITHIN_NUMERIC_RANGES = {
    24,  # X: supplier_name
    84,  # CF: receipt2_notice_no
}
NUMBER_COLUMNS = (
    set(range(18, 39))
    | {42, 43, 46, 49, 52, 53, 57, 60, 61, 62}
    | set(range(63, 68))
    | {70, 72}
    | set(range(76, 79))
    | set(range(81, 90))
    | {91, 92}
) - DATE_COLUMNS - PERCENT_COLUMNS - TEXT_COLUMNS_WITHIN_NUMERIC_RANGES


def _copy_header_alignment(header: Any, target: Any) -> None:
    target.alignment = copy(header.alignment)
    target.border = copy(header.border)


def _date_value(row: dict[str, Any], date_key: str, text_key: str | None = None) -> Any:
    value = row.get(date_key)
    if value is not None:
        return value
    return row.get(text_key) if text_key else None


def _row_values(row: dict[str, Any]) -> list[Any]:
    return [
        row.get("gross_net_type"), row.get("project_code"), row.get("department"), row.get("branch_company"),
        row.get("account_manager"), row.get("order_date"), row.get("business_type"), row.get("statistic_category"),
        row.get("team_level3_name"), row.get("customer_unit_name"), row.get("end_user_name"), row.get("regional_platform"),
        row.get("order_no"), row.get("project_name"), row.get("goods_name"), row.get("specification_model"),
        row.get("unit_name"), row.get("quantity"), _ratio(row.get("sales_tax_rate")),
        row.get("sales_unit_price_no_tax"), row.get("sales_unit_price"),
        row.get("revenue_no_tax"), row.get("order_value"), row.get("supplier_name"), _ratio(row.get("purchase_tax_rate")),
        row.get("purchase_unit_price_no_tax"),
        row.get("purchase_unit_price"), row.get("cost_no_tax"), row.get("purchase_amount"), row.get("delivery_date"),
        row.get("delivery_quantity"), row.get("delivery_revenue_no_tax"), row.get("delivery_value"),
        row.get("delivery_cost_no_tax"), row.get("delivery_cost"), row.get("pending_delivery_quantity"),
        row.get("pending_delivery_amount_no_tax"), row.get("pending_delivery_amount"), row.get("purchase_contract_no"),
        row.get("payment_terms"), row.get("purchase_performance_period"), row.get("purchase_signed_amount"),
        row.get("purchase_unsigned_amount"), _date_value(row, "received_invoice_date", "received_invoice_date_text"),
        row.get("purchase_invoice_no"), row.get("purchase_invoice_amount"),
        _date_value(row, "warehouse_date", "warehouse_date_text"), row.get("warehouse_voucher_no"),
        row.get("warehouse_amount"), _date_value(row, "booked_date", "booked_date_text"), row.get("booked_voucher_code"),
        row.get("booked_amount"), row.get("pending_booked_amount"), row.get("payment1_due_date"),
        _date_value(row, "payment1_date", "payment1_date_text"), row.get("payment1_voucher_no"), row.get("payment1_amount"),
        _date_value(row, "payment2_date", "payment2_date_text"), row.get("payment2_voucher_no"), row.get("payment2_amount"),
        row.get("total_paid"), row.get("accounts_payable"), row.get("gross_profit_no_tax"), row.get("tax_difference"), row.get("tax_refund"),
        row.get("gross_profit"), _ratio(row.get("gross_profit_margin_no_tax")),
        _date_value(row, "contract_signed_date", "contract_signed_date_text"), row.get("sales_contract_no"),
        row.get("sales_contract_value"), row.get("sales_performance_period"), row.get("sales_unsigned_contract_amount"),
        row.get("invoice_doc_no"), _date_value(row, "invoice_date", "invoice_date_text"), row.get("sales_invoice_no"),
        row.get("sales_invoice_amount"), row.get("pending_invoice_amount"), row.get("delivered_not_invoiced_amount"),
        _date_value(row, "receipt1_date", "receipt1_date_text"), row.get("receipt1_notice_no"), row.get("receipt1_amount"),
        _ratio(row.get("receipt1_ratio")), _date_value(row, "receipt2_date", "receipt2_date_text"),
        row.get("receipt2_notice_no"), row.get("receipt2_amount"), _ratio(row.get("receipt2_ratio")),
        row.get("total_received"), row.get("delivery_accounts_receivable"),
        row.get("invoice_accounts_receivable"), row.get("close_status"),
        row.get("labor_cost"), row.get("other_cost"),
    ]


def _ratio(value: Any) -> Any:
    if value is None:
        return None
    return Decimal(str(value)) / Decimal("100")


def _excel_column_name(column_no: int) -> str:
    value = column_no
    label = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        label = chr(65 + remainder) + label
    return label


def _export_sql(where_sql: str) -> str:
    return f"""
        SELECT
          ol.id AS order_line_id,
          so.gross_net_type, p.project_code, fin.department, fin.branch_company, fin.account_manager,
          fin.order_date, so.business_type, so.statistic_category, fin.team_level3_name,
          fin.customer_unit_name, fin.end_user_name, fin.regional_platform, so.order_no,
          COALESCE(ol.project_name, p.project_name) AS project_name,
          ol.goods_name, ol.specification_model, ol.unit_name, ol.quantity, ol.sales_tax_rate,
          ol.sales_unit_price_no_tax, ol.sales_unit_price, ol.revenue_no_tax, fin.profit_order_value AS order_value,
          pi.supplier_name, pi.purchase_tax_rate, pi.purchase_unit_price_no_tax, pi.purchase_unit_price, pi.cost_no_tax,
          fin.profit_purchase_amount AS purchase_amount, pi.labor_cost, pi.other_cost,
          fin.delivery_date, fin.delivery_quantity, fin.delivery_revenue_no_tax,
          fin.delivery_value, fin.delivery_cost_no_tax, fin.delivery_cost, fin.pending_delivery_quantity,
          fin.pending_delivery_amount_no_tax, fin.pending_delivery_amount,
          pc.purchase_contract_no, pc.payment_terms, pc.performance_period AS purchase_performance_period,
          pc.signed_amount AS purchase_signed_amount,
          COALESCE(fin.purchase_amount, 0) - COALESCE(fin.purchase_contract_signed_amount, 0) AS purchase_unsigned_amount,
          pinv.received_invoice_date, pinv.received_invoice_date_text, pinv.invoice_no AS purchase_invoice_no,
          pinv.invoice_amount AS purchase_invoice_amount, wh.warehouse_date, wh.warehouse_date_text,
          wh.voucher_no AS warehouse_voucher_no, wh.warehouse_amount,
          fpe.payment_date AS booked_date, fpe.payment_date_text AS booked_date_text,
          fpe.voucher_code AS booked_voucher_code, fpe.booked_amount,
          COALESCE(pi.purchase_amount, 0) - COALESCE(fpe_total.booked_amount, 0) AS pending_booked_amount,
          pay1.due_payment_date AS payment1_due_date, pay1.payment_date AS payment1_date,
          pay1.payment_date_text AS payment1_date_text, pay1.payment_voucher_no AS payment1_voucher_no,
          pay1.payment_amount AS payment1_amount, pay2.payment_date AS payment2_date,
          pay2.payment_date_text AS payment2_date_text, pay2.payment_voucher_no AS payment2_voucher_no,
          pay2.payment_amount AS payment2_amount, COALESCE(pay_total.payment_amount, 0) AS total_paid,
          COALESCE(pi.purchase_amount, 0) - COALESCE(pay_total.payment_amount, 0) AS accounts_payable,
          COALESCE(ol.revenue_no_tax, 0) - COALESCE(pi.cost_no_tax, 0) AS gross_profit_no_tax,
          fin.tax_difference, fin.tax_refund, fin.gross_profit,
          CASE WHEN COALESCE(ol.revenue_no_tax, 0) = 0 THEN 0
            ELSE (COALESCE(ol.revenue_no_tax, 0) - COALESCE(pi.cost_no_tax, 0)) / ol.revenue_no_tax * 100 END
            AS gross_profit_margin_no_tax,
          sc.contract_signed_date, sc.contract_signed_date_text, sc.sales_contract_no,
          sc.contract_value AS sales_contract_value, sc.performance_period AS sales_performance_period,
          sc.unsigned_contract_amount AS sales_unsigned_contract_amount,
          sinv.invoice_doc_no, sinv.invoice_date, sinv.invoice_date_text, sinv.invoice_no AS sales_invoice_no,
          sinv.invoice_amount AS sales_invoice_amount,
          COALESCE(fin.order_value, 0) - COALESCE(fin.sales_invoice_amount, 0) AS pending_invoice_amount,
          COALESCE(fin.delivery_value, 0) - COALESCE(fin.sales_invoice_amount, 0) AS delivered_not_invoiced_amount, rec1.receipt_date AS receipt1_date,
          rec1.receipt_date_text AS receipt1_date_text, rec1.payment_notice_no AS receipt1_notice_no,
          rec1.receipt_amount AS receipt1_amount,
          ROUND(rec1.receipt_amount / NULLIF(fin.sales_invoice_amount, 0) * 100, 6) AS receipt1_ratio,
          rec2.receipt_date AS receipt2_date, rec2.receipt_date_text AS receipt2_date_text,
          rec2.payment_notice_no AS receipt2_notice_no, rec2.receipt_amount AS receipt2_amount,
          ROUND(rec2.receipt_amount / NULLIF(fin.sales_invoice_amount, 0) * 100, 6) AS receipt2_ratio,
          COALESCE(rec_total.receipt_amount, 0) AS total_received,
          COALESCE(ol.order_value, 0) - COALESCE(rec_total.receipt_amount, 0) AS accounts_receivable,
          fin.delivery_accounts_receivable, fin.invoice_accounts_receivable,
          so.close_status
        FROM project p
        JOIN sales_order so ON so.project_id = p.id
        JOIN order_line ol ON ol.sales_order_id = so.id
        LEFT JOIN sub_project sp ON sp.id = ol.sub_project_id AND sp.deleted_at IS NULL
        LEFT JOIN v_order_line_finance fin ON fin.order_line_id = ol.id
        LEFT JOIN purchase_info pi ON pi.order_line_id = ol.id AND pi.deleted_at IS NULL
        LEFT JOIN purchase_contract pc ON pc.id = (
          SELECT MIN(pc0.id) FROM purchase_contract pc0 WHERE pc0.order_line_id = ol.id AND pc0.deleted_at IS NULL
        )
        LEFT JOIN purchase_invoice pinv ON pinv.order_line_id = ol.id AND pinv.phase_no = 1 AND pinv.deleted_at IS NULL
        LEFT JOIN warehouse_entry wh ON wh.order_line_id = ol.id AND wh.phase_no = 1 AND wh.deleted_at IS NULL
        LEFT JOIN finance_payment_entry fpe ON fpe.order_line_id = ol.id AND fpe.phase_no = 1 AND fpe.deleted_at IS NULL
        LEFT JOIN purchase_payment pay1 ON pay1.order_line_id = ol.id AND pay1.phase_no = 1 AND pay1.deleted_at IS NULL
        LEFT JOIN purchase_payment pay2 ON pay2.order_line_id = ol.id AND pay2.phase_no = 2 AND pay2.deleted_at IS NULL
        LEFT JOIN sales_contract sc ON sc.id = (
          SELECT MIN(sc0.id) FROM sales_contract sc0 WHERE sc0.order_line_id = ol.id AND sc0.deleted_at IS NULL
        )
        LEFT JOIN sales_invoice sinv ON sinv.order_line_id = ol.id AND sinv.phase_no = 1 AND sinv.deleted_at IS NULL
        LEFT JOIN sales_receipt rec1 ON rec1.order_line_id = ol.id AND rec1.phase_no = 1 AND rec1.deleted_at IS NULL
        LEFT JOIN sales_receipt rec2 ON rec2.order_line_id = ol.id AND rec2.phase_no = 2 AND rec2.deleted_at IS NULL
        LEFT JOIN (
          SELECT order_line_id, SUM(COALESCE(booked_amount, 0)) AS booked_amount
          FROM finance_payment_entry WHERE deleted_at IS NULL GROUP BY order_line_id
        ) fpe_total ON fpe_total.order_line_id = ol.id
        LEFT JOIN (
          SELECT order_line_id, SUM(COALESCE(payment_amount, 0)) AS payment_amount
          FROM purchase_payment WHERE deleted_at IS NULL GROUP BY order_line_id
        ) pay_total ON pay_total.order_line_id = ol.id
        LEFT JOIN (
          SELECT order_line_id, SUM(COALESCE(receipt_amount, 0)) AS receipt_amount
          FROM sales_receipt WHERE deleted_at IS NULL GROUP BY order_line_id
        ) rec_total ON rec_total.order_line_id = ol.id
        WHERE {where_sql}
        ORDER BY fin.order_date DESC, so.order_no, ol.id
    """


def filtered_export_rows(conn, user, filters=None):
    conditions = ["p.deleted_at IS NULL", "so.deleted_at IS NULL", "ol.deleted_at IS NULL"]
    params: dict[str, object] = {}
    active_filters = filters or {}
    project_text_filters = {
        "project_id": ("p.project_code", True),
        "department": ("fin.department", False),
        "client_unit": ("fin.customer_unit_name", True),
    }
    for key, (column, use_like) in project_text_filters.items():
        value = active_filters.get(key)
        if value in (None, ""):
            continue
        conditions.append(department_filter(column, key) if key == "department" else f"{column} {'LIKE' if use_like else '='} :{key}")
        params[key] = f"%{value}%" if use_like else value

    if active_filters.get("manager"):
        conditions.append(manager_match("p.project_code") if active_filters.get("include_history_manager") else "fin.account_manager LIKE :manager")
        params["manager"] = f"%{active_filters['manager']}%"

    # Detail-level filters must constrain the current exported row. Filtering only
    # by project existence would leak sibling orders from the same project.
    order_id = active_filters.get("order_id")
    if order_id not in (None, ""):
        conditions.append(order_match("ol.id"))
        params["order_id"] = f"%{order_id}%"
    start_date = active_filters.get("start_date")
    if start_date not in (None, ""):
        conditions.append("fin.order_date >= :start_date")
        params["start_date"] = start_date
    end_date = active_filters.get("end_date")
    if end_date not in (None, ""):
        conditions.append("fin.order_date <= :end_date")
        params["end_date"] = end_date

    order_status = active_filters.get("order_status")
    if order_status not in (None, ""):
        conditions.append(
            """
            EXISTS (
              SELECT 1
              FROM v_project_ledger_summary export_summary
              WHERE export_summary.project_code = p.project_code
                AND export_summary.computed_close_status = :order_status
            )
            """
        )
        params["order_status"] = order_status

    supplier_name = active_filters.get("supplier_name")
    if supplier_name not in (None, ""):
        conditions.append("pi.supplier_name LIKE :supplier_name")
        params["supplier_name"] = f"%{supplier_name}%"

    invoice_conditions = [
        "export_invoice.order_line_id = ol.id",
        "export_invoice.deleted_at IS NULL",
    ]
    invoice_start_date = active_filters.get("invoice_start_date")
    if invoice_start_date not in (None, ""):
        invoice_conditions.append("export_invoice.invoice_date >= :invoice_start_date")
        params["invoice_start_date"] = invoice_start_date
    invoice_end_date = active_filters.get("invoice_end_date")
    if invoice_end_date not in (None, ""):
        invoice_conditions.append("export_invoice.invoice_date <= :invoice_end_date")
        params["invoice_end_date"] = invoice_end_date
    if len(invoice_conditions) > 2:
        conditions.append(
            f"""
            EXISTS (
              SELECT 1
              FROM sales_invoice export_invoice
              WHERE {' AND '.join(invoice_conditions)}
            )
            """
        )

    apply_department_scope(conditions, params, user, "fin.department")
    return conn.execute(text(_export_sql(" AND ".join(conditions))), params).mappings().all()
