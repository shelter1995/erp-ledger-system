"""Explicit source-delivery allocation rules; totals survive rounding."""
from decimal import Decimal, ROUND_DOWN, ROUND_HALF_UP

from openpyxl.utils import get_column_letter

from .legacy_ledger_parser import parse_date_sequence, parse_amount_sequence
from .validation import validate_business_date

DELIVERY_COLUMNS = {
    31: 'delivery_quantity', 32: 'delivery_revenue_no_tax', 33: 'delivery_value',
    34: 'delivery_cost_no_tax', 35: 'delivery_cost', 36: 'pending_delivery_quantity',
    37: 'pending_delivery_amount_no_tax', 38: 'pending_delivery_amount',
}


def preserve_multiple_deliveries(conn, order_line_id, data):
    """Reject aggregate edits; unchanged form saves must retain every phase."""
    from fastapi import HTTPException
    from sqlalchemy import text
    from .financial_calculations import line_snapshot

    count = conn.execute(text('SELECT COUNT(*) FROM delivery_record WHERE order_line_id=:id AND deleted_at IS NULL'), {'id':order_line_id}).scalar_one()
    if count <= 1:
        return False
    current = line_snapshot(conn, order_line_id)
    if any(value != current.get(key) for key, value in data.items()):
        raise HTTPException(409, '该明细有多笔交付记录，不能通过合计字段覆盖；请核对各笔交付记录')
    return True


def delivery_phases(row, row_no):
    parsed = parse_date_sequence(row[29])
    if parsed.blocking:
        raise ValueError(f'AD{row_no}：交付日期无法解析，请核对')
    dates = [validate_business_date(d) for d in parsed.values] or [None]
    count = len(dates)
    phases = [{'delivery_date': d} for d in dates]
    allocated = False
    for col, field in DELIVERY_COLUMNS.items():
        amount = parse_amount_sequence(row[col - 1])
        if amount.blocking or any(not v.is_finite() or abs(v) > Decimal('999999999999.99') for v in amount.values):
            raise ValueError(f'{get_column_letter(col)}{row_no}：交付数值无法解析或超出范围')
        scale = Decimal('.000001') if col in (31, 36) else Decimal('.01')
        values = [v.quantize(scale, rounding=ROUND_HALF_UP) for v in amount.values]
        if not values:
            values = [None] * count
        elif len(values) == 1 and count > 1:
            total = values[0]
            if col >= 36:
                # Undelivered figures are the remaining line balance, not a
                # further delivery. Retain it once on the final delivery.
                values = [None] * (count - 1) + [total]
            else:
                share = (total / count).quantize(scale, rounding=ROUND_DOWN)
                values = [share] * (count - 1) + [total - share * (count - 1)]
                allocated = True
        elif len(values) != count:
            raise ValueError(f'{get_column_letter(col)}{row_no}：交付日期与数值笔数不一致，须逐笔对应或只填一个合计')
        for phase, value in zip(phases, values):
            phase[field] = value
    return phases, allocated
