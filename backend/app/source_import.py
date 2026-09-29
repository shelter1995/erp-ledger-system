"""Import an authoritative ledger using confirmed positional allocation rules.

Source rows are business records, including repeated rows. File digests prevent
replaying a completed file. All mutations run in the caller's locked transaction.
"""
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP, ROUND_DOWN
from hashlib import sha256
from io import BytesIO
import json
import posixpath
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from sqlalchemy import bindparam, text

from .importer import import_excel, validate_template_numbers
from .ledger_excel import (
    LEGACY_TEMPLATE_HEADERS,
    NUMBER_COLUMNS,
    PERCENT_COLUMNS,
    TEMPLATE_HEADERS,
    normalize_standard_row,
    standard_template_version,
)
from .legacy_ledger_parser import (
    merge_order_number_history,
    parse_amount_sequence,
    parse_date_sequence,
    parse_document_sequence,
)
from .legacy_import_service import FINANCE_SOURCES, PHASE_TABLE_COLUMNS
from .validation import validate_business_date
from .financial_calculations import refresh_balances_many, update_fields_many
from .edit_versions import touch_lines
from .source_delivery import delivery_phases, DELIVERY_COLUMNS
from .historical_project_identity import (
    ProjectResolution,
    _is_generated_temporary_code,
    resolve_project_code,
)


TRANSITIONAL_91_HEADERS = (
    LEGACY_TEMPLATE_HEADERS[:24]
    + LEGACY_TEMPLATE_HEADERS[25:87]
    + TEMPLATE_HEADERS[87:92]
)
TRANSITIONAL_92_HEADERS = LEGACY_TEMPLATE_HEADERS[:87] + TEMPLATE_HEADERS[87:92]
MAX_ORDER_NUMBER_LENGTH = 64
SOURCE_IMPORT_HISTORY_SOURCE = 'source_import'


def source_template_layout(headers):
    names = [str(value or '').strip() for value in headers]
    while names and not names[-1]:
        names.pop()
    version = standard_template_version(names)
    if version:
        return version
    if names == TRANSITIONAL_91_HEADERS:
        return 'transitional_91'
    if names == TRANSITIONAL_92_HEADERS:
        return 'transitional_92'
    return None


def normalize_source_row(values, layout):
    if layout == 'transitional_91':
        row = list(values[:91]) + [None] * max(0, 91 - len(values))
        row.insert(24, None)
        return row
    if layout == 'transitional_92':
        return list(values[:92]) + [None] * max(0, 92 - len(values))
    return normalize_standard_row(values, layout)


def source_column(layout, standard_column):
    if layout == 91 and standard_column >= 89:
        return None if standard_column == 89 else standard_column - 1
    if layout != 'transitional_91':
        return standard_column
    if standard_column == 25:
        return None
    return standard_column if standard_column < 25 else standard_column - 1


def restore_source_number_formats(content, ws, layout):
    """Recover numeric XML values misread as dates in known non-date fields.

    Read the stored number, never reverse-convert a datetime (which has already
    lost precision or become #VALUE!). Real date columns and literal Excel
    errors are untouched. Only the in-memory normalization copy is changed.
    """
    document_columns = {45, 48, 51, 56, 59, 73, 75, 80, 84}
    columns = {
        source_column(layout, col): col
        for col in NUMBER_COLUMNS | PERCENT_COLUMNS | document_columns
        if source_column(layout, col) is not None
    }
    candidates = {
        cell.coordinate: columns[cell.column]
        for row in ws.iter_rows(min_row=3)
        for cell in row
        if cell.column in columns and cell.value is not None
        and is_date_format(cell.number_format)
    }
    if not candidates:
        return []
    main_ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    rel_ns = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    warnings = []
    with ZipFile(BytesIO(content)) as archive:
        workbook = ElementTree.fromstring(archive.read('xl/workbook.xml'))
        sheet = next(s for s in workbook.find(main_ns + 'sheets') if s.get('name') == ws.title)
        relationships = ElementTree.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        relation = next(r for r in relationships if r.get('Id') == sheet.get(rel_ns + 'id'))
        target = relation.get('Target')
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
        with archive.open(path) as source:
            for _, node in ElementTree.iterparse(source, events=('end',)):
                if node.tag == main_ns + 'c':
                    coordinate = node.get('r')
                    value = node.find(main_ns + 'v')
                    if coordinate in candidates and node.get('t', 'n') == 'n' and value is not None and value.text:
                        col = candidates.pop(coordinate)
                        cell = ws[coordinate]
                        # Match openpyxl's numeric conversion; document numbers
                        # remain text so their digits survive the next workbook read.
                        cell.value = (format(Decimal(value.text), 'f') if col in document_columns
                                      else float(value.text) if any(c in value.text for c in '.eE')
                                      else int(value.text))
                        cell.number_format = 'General'
                        warnings.append({
                            'row': cell.row, 'field': TEMPLATE_HEADERS[col - 1],
                            'message': f'{coordinate} 误设为日期格式，已按原文件保存的数值读取，请核对；原文件未修改',
                        })
                    node.clear()
                elif node.tag == main_ns + 'row':
                    node.clear()
                if not candidates:
                    break
    return warnings


def source_text(value):
    if value is None:
        return None
    return value.isoformat() if isinstance(value, (datetime, date)) else str(value)


def source_phases(date_value, amount_value, document_value, *, split_single_total=False):
    """Only split proven positional pairs. Otherwise retain one undated total."""
    dates = parse_date_sequence(date_value)
    amounts = parse_amount_sequence(amount_value)
    documents = parse_document_sequence(document_value)
    if amounts.blocking or any(not a.is_finite() or abs(a) > Decimal('9999999999999999.99') for a in amounts.values):
        raise ValueError('财务金额无效或超出范围')
    total = sum(amounts.values, Decimal(0)) if amounts.values else None
    if total is not None:
        total = total.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    if (split_single_total and not dates.blocking and len(dates.values) > 1
            and len(amounts.values) in (1, len(dates.values))
            and len(documents.values) in (0, 1, len(dates.values))):
        count = len(dates.values)
        if len(amounts.values) == 1:
            share = (total / count).quantize(Decimal('.01'), rounding=ROUND_DOWN)
            values = [share] * (count - 1) + [total - share * (count - 1)]
        else:
            values = [v.quantize(Decimal('.01'), rounding=ROUND_HALF_UP) for v in amounts.values]
        refs = (documents.values * count if len(documents.values) == 1
                else documents.values or [None] * count)
        return [(validate_business_date(d), d.isoformat(), value, ref)
                for d, value, ref in zip(dates.values, values, refs)], False
    if len(dates.values) == len(amounts.values) > 1 and not dates.blocking and len(documents.values) in (0, len(dates.values)):
        return [(validate_business_date(d), d.isoformat(), a.quantize(Decimal('.01'), rounding=ROUND_HALF_UP), documents.values[i] if documents.values else None)
                for i, (d, a) in enumerate(zip(dates.values, amounts.values))], False
    actual_date = validate_business_date(dates.values[0]) if not dates.blocking and len(dates.values) == 1 else None
    return [(actual_date, source_text(date_value), total, source_text(document_value))], bool(
        dates.blocking or len(dates.values) > 1 or (total is not None and actual_date is None))


def source_order_chain(value, row_no):
    normalized = re.sub(r'[/／](?:\s*[/／])+', '/', source_text(value) or '')
    parsed = merge_order_number_history([normalized], order_key=normalized)
    if parsed.blocking:
        detail = '；'.join(issue.message for issue in parsed.issues if issue.blocking)
        raise ValueError(f'第 {row_no} 行销售订单号无效：{detail}')
    chain = [str(item).strip() for item in parsed.values if str(item).strip()]
    if not chain:
        raise ValueError(f'第 {row_no} 行缺少销售订单号')
    overlong = next((item for item in chain if len(item) > MAX_ORDER_NUMBER_LENGTH), None)
    if overlong:
        raise ValueError(
            f'第 {row_no} 行销售订单号中的单个编号“{overlong}”超过 '
            f'{MAX_ORDER_NUMBER_LENGTH} 个字符'
        )
    return chain


def _project_resolution_state(conn, order_groups):
    """Prefetch the two lookups otherwise repeated for every source order."""
    order_nos = sorted(order_groups)
    official_codes = sorted({
        str(row[1]).strip()
        for grouped_rows in order_groups.values()
        for _, row in grouped_rows
        if str(row[1] or '').strip()
    })
    projects_by_order = {}
    for start in range(0, len(order_nos), 1000):
        rows = conn.execute(
            text('''SELECT DISTINCT so.order_no,p.project_code
                FROM sales_order so JOIN project p ON p.id=so.project_id
                WHERE so.order_no IN :order_nos AND so.deleted_at IS NULL
                AND p.deleted_at IS NULL''').bindparams(bindparam('order_nos', expanding=True)),
            {'order_nos': order_nos[start:start + 1000]},
        ).mappings()
        for row in rows:
            projects_by_order.setdefault(str(row['order_no']), []).append(str(row['project_code']))
    deleted_official_codes = set()
    for start in range(0, len(official_codes), 1000):
        rows = conn.execute(
            text('SELECT project_code,deleted_at FROM project WHERE project_code IN :codes')
            .bindparams(bindparam('codes', expanding=True)),
            {'codes': official_codes[start:start + 1000]},
        ).mappings()
        deleted_official_codes.update(
            str(row['project_code']) for row in rows if row['deleted_at'] is not None
        )
    return projects_by_order, deleted_official_codes


def _has_other_temporary_project(projects_by_order, order_no, official_code):
    return any(
        code != official_code and _is_generated_temporary_code(code, order_no)
        for code in projects_by_order.get(order_no, [])
    )


def persist_order_histories(conn, line_ids, order_chains):
    """Index all source aliases without inventing one linear history per order.

    Rows may converge on one current number or diverge from a shared old number.
    Their exact ordered paths remain in ledger_raw_row and are exposed per line.
    Existing orders/lines are never reassigned by importing another source path.
    """
    if not line_ids:
        return
    owners = {int(row['order_line_id']): row for row in conn.execute(
        text('''SELECT ol.id AS order_line_id,so.id AS sales_order_id,so.order_no
            FROM order_line ol JOIN sales_order so ON so.id=ol.sales_order_id
            WHERE ol.id IN :ids''').bindparams(bindparam('ids', expanding=True)),
        {'ids': list(line_ids.values())},
    ).mappings()}
    targets = {}
    for row_no, line_id in line_ids.items():
        owner = owners[int(line_id)]
        target = targets.setdefault(int(owner['sales_order_id']), {'current':owner['order_no'], 'aliases':[]})
        target['aliases'].extend(order_chains[row_no])
    histories = {}
    for row in conn.execute(text('''SELECT sales_order_id,order_no FROM sales_order_number_history
        WHERE sales_order_id IN :ids ORDER BY history_order''').bindparams(bindparam('ids', expanding=True)),
        {'ids':list(targets)}).mappings():
        histories.setdefault(int(row['sales_order_id']), []).append(row['order_no'])
    inserts = []
    for order_id, target in targets.items():
        aliases = list(dict.fromkeys(histories.get(order_id, []) + target['aliases']))
        aliases = [number for number in aliases if number != target['current']] + [target['current']]
        inserts.extend({'sales_order_id':order_id, 'order_no':number, 'history_order':position,
                        'source':SOURCE_IMPORT_HISTORY_SOURCE}
                       for position,number in enumerate(aliases,1))
    statement = text('''INSERT INTO sales_order_number_history
        (sales_order_id,order_no,history_order,source)
        VALUES (:sales_order_id,:order_no,:history_order,:source)
        ON DUPLICATE KEY UPDATE order_no=VALUES(order_no),source=VALUES(source)''')
    for start in range(0,len(inserts),1000):
        conn.execute(statement,inserts[start:start+1000])


def import_source(conn, content, filename, user, *, preview=False, duplicate_confirmation=None):
    digest = sha256(content).hexdigest()
    existing = conn.execute(text("SELECT id,success_rows FROM import_batch WHERE source_sha256=:sha AND status IN ('completed','reverted') LIMIT 1"), {'sha':digest}).mappings().first()
    if existing:
        raise HTTPException(409, f"这份文件已经导入（批次 {existing['id']}，{existing['success_rows']} 行），未重复写入")
    wb = load_workbook(BytesIO(content), data_only=True)
    candidates = [s for s in wb if source_template_layout([c.value for c in s[2]])]
    if len(candidates) != 1:
        raise ValueError('请使用包含唯一业务工作表的 0916 新版92列或原版91列模板')
    ws = candidates[0]
    version = source_template_layout([c.value for c in ws[2]])
    originals, records, deliveries = {}, {}, {}
    warnings = restore_source_number_formats(content, ws, version)
    groups = dict(FINANCE_SOURCES)
    groups['finance_payment_entry'] = groups.pop('finance_invoice_check')
    # Sales document number and invoice number are separate source columns.
    groups['sales_invoice'] = ((74, 76, 75),)
    prepared_rows = []
    order_chains = {}
    order_groups = {}
    for row_no, cells in enumerate(ws.iter_rows(min_row=3), 3):
        values = [c.value for c in cells]
        if not any(v not in (None, '') for v in values):
            continue
        if len(originals) >= 20000:
            raise ValueError('单次最多20,000行')
        row = normalize_source_row(values, version)
        if not row[12] or not row[14]:
            raise ValueError(f'第 {row_no} 行缺少销售订单号或物资/服务名称')
        order_chain = source_order_chain(row[12], row_no)
        if re.search(r'[/／]\s*[/／]', source_text(row[12]) or ''):
            warnings.append({'row':row_no, 'field':'销售订单号', 'message':'连续斜杠按一个分隔符解析，保留已有历史编号顺序'})
        originals[row_no] = row[:]
        order_chains[row_no] = order_chain
        order_no = order_chain[-1]
        prepared_rows.append((row_no, cells, row))
        order_groups.setdefault(order_no, []).append((row_no, row))

    projects_by_order, deleted_official_codes = _project_resolution_state(conn, order_groups)
    reserved_codes = set()
    effective_codes = {}
    for order_no, grouped_rows in order_groups.items():
        official_codes = {str(row[1]).strip() for _, row in grouped_rows if str(row[1] or '').strip()}
        if len(official_codes) > 1:
            blank_rows = [row_no for row_no, row in grouped_rows if not str(row[1] or '').strip()]
            if blank_rows:
                raise ValueError(
                    f'订单号 {order_no} 在同一文件中填写了多个项目编号，'
                    f'第 {blank_rows[0]} 行 B 列为空，无法判断应归入哪个项目'
                )
            for official_code in sorted(official_codes):
                official_rows = [(row_no, row) for row_no, row in grouped_rows
                                 if str(row[1]).strip() == official_code]
                first_row_no, first_row = official_rows[0]
                resolution = (
                    resolve_project_code(
                        conn,
                        project_code=official_code,
                        order_no=order_no,
                        first_goods_name=first_row[14],
                        reserved_codes=reserved_codes,
                        reconcile_temporary=False,
                    )
                    if _has_other_temporary_project(
                        projects_by_order, order_no, official_code
                    )
                    else ProjectResolution(official_code)
                )
                for row_no, _ in official_rows:
                    effective_codes[row_no] = resolution.project_code
            continue
        first_row_no, first_row = grouped_rows[0]
        official_code = next(iter(official_codes), None)
        if official_code in deleted_official_codes:
            raise ValueError(f'正式项目编号 {official_code} 已被停用，不能自动恢复或合并')
        resolution = (
            resolve_project_code(
                conn,
                project_code=official_code,
                order_no=order_no,
                first_goods_name=first_row[14],
                reserved_codes=reserved_codes,
            )
            if official_code is None or _has_other_temporary_project(
                projects_by_order, order_no, official_code
            )
            else ProjectResolution(official_code)
        )
        for row_no, _ in grouped_rows:
            effective_codes[row_no] = resolution.project_code
        if resolution.message:
            warnings.append({'row':first_row_no, 'field':'项目编号', 'message':resolution.message})

    for row_no, cells, row in prepared_rows:
        row[1] = effective_codes[row_no]
        row[12] = order_chains[row_no][-1]
        deliveries[row_no], allocated = delivery_phases(row, row_no)
        if allocated:
            warnings.append({'row':row_no, 'field':'交付情况', 'message':'多日期对应单个合计时按期均分金额和数量；尾差计入最后一期，合计不变'})
        # Feed scalar totals to the base importer; replace its provisional
        # delivery row with the proven positional phases after import succeeds.
        row[29] = deliveries[row_no][0]['delivery_date']
        for col, field in DELIVERY_COLUMNS.items():
            values = [p[field] for p in deliveries[row_no] if p[field] is not None]
            row[col - 1] = sum(values, Decimal(0)) if values else None
        for col in (19,25):
            original_col = source_column(version, col)
            if original_col and cells[original_col-1].data_type == 'e':
                warnings.append({'row':row_no,'field':TEMPLATE_HEADERS[col-1],'message':f'原表税率公式为 {row[col-1]}，保留原始值，税率留空，不影响已填金额'})
                row[col-1] = None
        records[row_no] = {}
        for name, sources in groups.items():
            phases = []
            for dc, ac, nc in sources:
                if all(row[c-1] in (None, '') for c in (dc, ac, nc)):
                    continue
                try:
                    parts, aggregate = source_phases(row[dc-1], row[ac-1], row[nc-1], split_single_total=name == 'purchase_invoice')
                except ValueError as exc:
                    coordinate = f'{get_column_letter(source_column(version, ac))}{row_no}'
                    raise ValueError(f'{coordinate}（第 {row_no} 行，{TEMPLATE_HEADERS[ac-1]}）：{exc}') from exc
                phases.extend(parts)
                if name == 'purchase_invoice' and len(parts) > 1 and len(parse_amount_sequence(row[ac-1]).values) == 1:
                    warnings.append({'row':row_no, 'field':TEMPLATE_HEADERS[ac-1],
                                     'message':'多收票日期对应单个合计，按期均分，尾差计入最后一期；单一发票号沿用原值'})
                if aggregate:
                    warnings.append({'row':row_no, 'field':TEMPLATE_HEADERS[ac-1], 'message':'保留原日期文本及合计金额，不分摊；无单一日期的金额不归入某一天'})
            records[row_no][name] = phases
        for sources in groups.values():
            for columns in sources:
                for col in columns:
                    row[col-1] = None
        # Prevent the normalized base import from creating a due-date-only
        # payment that this authoritative phase import would immediately delete.
        row[53] = None
        row[72] = None  # invoice_doc_no is restored with the sales invoice records
        # Excel binary precision is normalized to the database's documented scale.
        # Original unrounded values remain in the positional source archive.
        for col in (18,20,21,26,27,31,23,29,42,70,91,92):
            value = row[col-1]
            if value in (None, ''):
                continue
            number = Decimal(str(value))
            if not number.is_finite() or abs(number) > Decimal('999999999999.99'):
                raise ValueError(f'第 {row_no} 行 {TEMPLATE_HEADERS[col-1]} 无效或超出范围')
            scale = Decimal('.000001') if col in (18,20,21,26,27,31) else Decimal('.01')
            row[col-1] = number.quantize(scale, rounding=ROUND_HALF_UP)
        for col, value in enumerate(row, 1):
            ws.cell(row_no, col).value = value
    for col, label in enumerate(TEMPLATE_HEADERS, 1):
        ws.cell(2,col).value = label
    for other in list(wb):
        if other != ws:
            wb.remove(other)
    if not originals:
        raise ValueError('文件没有业务数据')
    from .import_review import find_duplicates, summary
    duplicates = find_duplicates(conn, originals, user, digest)
    if duplicates['count'] and not preview and duplicate_confirmation != duplicates['confirmation_token']:
        raise HTTPException(409, '发现疑似重复明细或比对结果已变化，请重新预检并确认保留后再导入')
    normalized = BytesIO(); wb.save(normalized); wb.close()
    source_archives = {}
    for row_no, row in originals.items():
        raw = json.dumps(
            {'headers': TEMPLATE_HEADERS, 'values': row},
            ensure_ascii=False,
            default=source_text,
        )
        source_archives[row_no] = {
            'raw_json': raw,
            'row_hash': sha256((str(row_no) + ':' + raw).encode()).hexdigest(),
        }
    ids = {}
    result = import_excel(conn, reset=False, workbook_bytes=normalized.getvalue(), source_file_name=filename,
                          user=user, strict_template=True, preserve_source=True,
                          line_id_map=ids, source_archives=source_archives)
    if result['failed_rows']:
        from .import_report import ImportReportError
        raise ImportReportError(result)
    persist_order_histories(conn, ids, order_chains)
    written = Counter()
    phase_inserts = {}
    authoritative = {
        'order_line': [],
        'purchase_info': [],
    }
    for row_no, line_id in ids.items():
        row = originals[row_no]
        for table, mapping in {
            'order_line': {
                'sales_unit_price_no_tax': 20,
                'sales_unit_price': 21,
                'revenue_no_tax': 22,
                'order_value': 23,
            },
            'purchase_info': {
                'purchase_unit_price_no_tax': 26,
                'purchase_unit_price': 27,
                'cost_no_tax': 28,
                'purchase_amount': 29,
            },
        }.items():
            values = {'id': line_id}
            for field, col in mapping.items():
                value = row[col - 1]
                values[field] = None if value in (None, '') else Decimal(str(value)).quantize(
                    Decimal('.000001')
                    if 'unit_price' in field or field.endswith('quantity')
                    else Decimal('.01'),
                    rounding=ROUND_HALF_UP,
                )
            authoritative[table].append(values)
        for name, phases in records[row_no].items():
            for no, (dt, dt_text, amount, document) in enumerate(phases, 1):
                params = {
                    'id': line_id,
                    'phase': no,
                    'date': dt,
                    'date_text': dt_text,
                    'amount': amount,
                    'document': document,
                }
                if name == 'sales_invoice':
                    params['invoice_doc_no'] = source_text(row[72])
                elif name == 'purchase_payment':
                    if no == 1 and row[53] is not None:
                        from .importer import _as_date
                        params['due_payment_date'] = _as_date(row[53])
                    else:
                        params['due_payment_date'] = None
                phase_inserts.setdefault(name, []).append(params)
                written[name] += 1
    update_fields_many(
        conn,
        'order_line',
        'id',
        authoritative['order_line'],
        ('sales_unit_price_no_tax', 'sales_unit_price', 'revenue_no_tax', 'order_value'),
    )
    update_fields_many(
        conn,
        'purchase_info',
        'order_line_id',
        authoritative['purchase_info'],
        ('purchase_unit_price_no_tax', 'purchase_unit_price', 'cost_no_tax', 'purchase_amount'),
    )
    delivery_rows = [dict(phase, order_line_id=ids[row_no])
                     for row_no, phases in deliveries.items() for phase in phases]
    line_ids = list(ids.values())
    for start in range(0, len(line_ids), 1000):
        conn.execute(text('DELETE FROM delivery_record WHERE order_line_id IN :ids')
                     .bindparams(bindparam('ids', expanding=True)), {'ids':line_ids[start:start+1000]})
    fields = ['order_line_id', 'delivery_date', *DELIVERY_COLUMNS.values()]
    statement = text(f'INSERT INTO delivery_record ({",".join(fields)}) VALUES ({",".join(":"+f for f in fields)})')
    for start in range(0, len(delivery_rows), 1000):
        conn.execute(statement, delivery_rows[start:start+1000])
    written['delivery_record'] = len(delivery_rows)
    for name, rows in phase_inserts.items():
        dc, nc, ac = PHASE_TABLE_COLUMNS[name]
        extra_column = ',invoice_doc_no' if name == 'sales_invoice' else (
            ',due_payment_date' if name == 'purchase_payment' else ''
        )
        extra_value = ',:invoice_doc_no' if name == 'sales_invoice' else (
            ',:due_payment_date' if name == 'purchase_payment' else ''
        )
        statement = text(
            f'INSERT INTO {name} '
            f'(order_line_id,phase_no,{dc},{dc}_text,{ac},{nc}{extra_column}) '
            f'VALUES (:id,:phase,:date,:date_text,:amount,:document{extra_value})'
        )
        for start in range(0, len(rows), 1000):
            conn.execute(statement, rows[start:start + 1000])
    # Preview is rolled back and its response reads the live financial view, so
    # stored compatibility fields and project versions do not need a trial write.
    # A real import performs the same finalization in bounded batches.
    if not preview:
        touch_lines(conn, ids.values())
        refresh_balances_many(conn, ids.values())
    conn.execute(text('UPDATE import_batch SET source_sha256=:sha WHERE id=:id'), {'sha':digest,'id':result['batch_id']})
    return {**result,'source_sha256':digest,'warnings':warnings,'phase_counts':dict(written),'source_preserved':True,
            'layout':{
                'transitional_91':'过渡版台账',
                'transitional_92':'过渡版台账（含采购税率）',
                92:'0916 新版',
                91:'原版台账',
            }[version],
            'summary':summary(conn,result['batch_id']), 'duplicates':duplicates}
