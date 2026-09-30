from io import BytesIO
import json
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from app.db import db
from test_financial_integration import _create_order, _basic_payload
from test_permission_routes_v2 import actor, directory


def transfer(client, headers, line, *, scope='order', **changes):
    response = client.get(f'/api/history/lines/{line}', headers=headers)
    assert response.status_code == 200, response.text
    current = response.json()['current']
    fields = ('account_manager', 'department', 'branch_company', 'team_level3_name')
    values = {k: current[k] for k in fields}
    return client.post(f'/api/history/transfer-{scope}', headers=headers,
                       json={'order_line_id':line, 'expected':values, **values,
                             **changes, 'reason':'单订单归属回归'})


@pytest.mark.parametrize('source', [False, True])
def test_order_transfer_isolates_siblings_history_names_and_export(client, headers, source):
    a, payload = _create_order(client, headers, 'OWN-A', department='甲部')
    b, other = _create_order(client, headers, 'OWN-B', project_code=payload['project_code'], department='甲部')
    a2, _ = _create_order(client, headers, 'OWN-A2', project_code=payload['project_code'],
                          order_no=payload['order_no'], goods_name='第二条明细', department='甲部')
    if source:
        with db() as conn:
            conn.execute(text('UPDATE order_line SET source_preserved=1, line_department=:d, '
                              'line_account_manager=:m, line_branch_company=:b, line_team_level3_name=:t'),
                         {'d':payload['department'], 'm':payload['account_manager'],
                          'b':payload['branch_company'], 't':payload['team_name']})
    response = transfer(client, headers, a, account_manager='新负责人', department='乙部')
    assert response.status_code == 200, response.text
    assert response.json()['affected_orders'] == 1
    assert response.json()['affected_lines'] == 2
    for line in (a, a2):
        ctx = client.get(f'/api/history/lines/{line}', headers=headers).json()
        assert ctx['current']['account_manager'] == '新负责人'
        assert ctx['current']['department'] == '乙部'
        assert [h['manager_name'] for h in ctx['managers']] == [payload['account_manager'], '新负责人']
    sibling = client.get(f'/api/history/lines/{b}', headers=headers).json()
    assert sibling['current']['department'] == '甲部'
    assert sibling['current']['account_manager'] == other['account_manager']
    assert '新负责人' not in str(sibling['managers'])
    for endpoint in ('sales', 'purchases'):
        for history in ('false', 'true'):
            result = client.get(f'/api/{endpoint}', headers=headers,
                                params={'manager':'新负责人', 'include_history_manager':history})
            assert result.status_code == 200, result.text
            assert {r['order_line_id'] for r in result.json()['items']} == {a, a2}
    rename = client.post('/api/history/rename-orders', headers=headers, json={'items':[
        {'order_line_id':a, 'expected_order_no':payload['order_no'], 'order_no':'ONLY-A-RENAMED', 'reason':'改当前订单号'}]})
    assert rename.status_code == 200, rename.text
    assert client.get(f'/api/history/lines/{b}', headers=headers).json()['current']['order_no'] == other['order_no']
    # Ordinary edits after an override cannot write the new owner back to project.
    edited = {**payload, 'order_no':'ONLY-A-RENAMED', 'account_manager':'新负责人',
              'department':'乙部', 'project_name':'独立项目名称', 'goods_name':'修改货物名称'}
    result = client.put(f'/api/orders/{a}', headers=headers, json=edited)
    assert result.status_code == 200, result.text
    assert client.get(f'/api/history/lines/{b}', headers=headers).json()['current']['department'] == '甲部'
    exported = client.get('/api/history/export', headers=headers, params={'manager':'新负责人','include_history_manager':'true'})
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(BytesIO(exported.content))
    assert workbook['明细'].max_row == 3
    assert '新负责人' in str(list(workbook['负责人历史'].values))
    assert other['order_no'] not in str(list(workbook['负责人历史'].values))
    with db() as conn:
        if source:
            assert conn.execute(text('SELECT line_department FROM order_line WHERE id=:id'), {'id':a}).scalar() == '甲部'
    # Later framework transfer includes both orders and preserves A's private history.
    response = transfer(client, headers, b, scope='project', account_manager='框架负责人')
    assert response.status_code == 200, response.text
    assert response.json()['affected_orders'] == 2
    history_a = client.get(f'/api/history/lines/{a}', headers=headers).json()['managers']
    history_b = client.get(f'/api/history/lines/{b}', headers=headers).json()['managers']
    assert '新负责人' in str(history_a)
    assert '新负责人' not in str(history_b)


def test_order_transfer_requires_permission_and_moves_department_visibility(client, headers):
    a, payload = _create_order(client, headers, 'SCOPE-A', department='甲部')
    b, _ = _create_order(client, headers, 'SCOPE-B', project_code=payload['project_code'], department='甲部')
    permissions = ['order_view','order_edit','sales_view','purchase_view','ledger_view']
    with db() as conn:
        da, db_id = directory(conn, '甲部'), directory(conn, '乙部')
        _, plain = actor(conn, 'order_transfer_plain', permissions, [da,db_id])
        _, limited = actor(conn, 'order_transfer_limited', permissions+['department_transfer'], [da])
        _, both = actor(conn, 'order_transfer_both', permissions+['department_transfer'], [da,db_id])
        _, target = actor(conn, 'order_transfer_target', permissions, [db_id])
    assert transfer(client, plain, a, account_manager='拒绝变更').status_code == 403
    assert transfer(client, limited, a, department='乙部').status_code == 403
    assert transfer(client, both, a, department='乙部', account_manager='乙部负责人').status_code == 200
    for endpoint in ('orders','sales','purchases'):
        old = client.get(f'/api/{endpoint}', headers=limited)
        new = client.get(f'/api/{endpoint}', headers=target)
        assert old.status_code == new.status_code == 200
        assert {r['order_line_id'] for r in old.json()['items']} == {b}
        assert {r['order_line_id'] for r in new.json()['items']} == {a}
    for endpoint in ('history/lines','sales','purchases'):
        assert client.get(f'/api/{endpoint}/{a}', headers=limited).status_code == 404
    assert client.get(f'/api/history/lines/{b}', headers=limited).json()['project_scope'] is None
    assert transfer(client, limited, b, scope='project', account_manager='越权整框架').status_code == 404
    assert transfer(client, limited, b, account_manager='只改乙订单').status_code == 200
    with db() as conn:
        log = conn.execute(text("SELECT department_ids_json FROM operation_log WHERE action_name='transfer_order' ORDER BY id LIMIT 1")).scalar()
        assert set(json.loads(log)) == {da,db_id}


def test_add_line_after_order_transfer_inherits_order_owner(client, headers):
    a, payload = _create_order(client, headers, 'ADD-A', department='甲部')
    b, other = _create_order(client, headers, 'ADD-B', project_code=payload['project_code'], department='甲部')
    assert transfer(client, headers, a, department='乙部', account_manager='新负责人').status_code == 200
    response = client.post('/api/orders', headers=headers, json={**payload, 'department':'乙部',
                           'account_manager':'新负责人', 'goods_name':'交接后新增明细'})
    assert response.status_code == 200, response.text
    new_id = response.json()['order_line_id']
    assert client.get(f'/api/history/lines/{new_id}', headers=headers).json()['current']['department'] == '乙部'
    assert client.get(f'/api/history/lines/{b}', headers=headers).json()['current']['department'] == '甲部'
    batch = client.post('/api/orders/batch-basic', headers=headers, json={'items':[
        {**_basic_payload(payload), 'department':'乙部', 'account_manager':'新负责人', 'goods_name':'批量新增甲'},
        {**_basic_payload(other), 'goods_name':'批量新增乙'}]})
    assert batch.status_code == 200, batch.text


def test_mixed_order_transfer_checks_every_original_department(client, headers):
    from test_permission_routes_v2 import seed_mixed
    with db() as conn:
        da, _, lines = seed_mixed(conn)
        _, limited = actor(conn, 'mixed_order_transfer', ['order_view','order_edit','department_transfer'], [da])
    ctx = client.get(f'/api/history/lines/{lines[0]}', headers=limited).json()
    assert ctx['order_scope'] is None
    result = transfer(client, limited, lines[0], account_manager='不得部分交接')
    assert result.status_code == 404, result.text
    with db() as conn:
        assert conn.execute(text('SELECT SUM(ownership_overridden) FROM sales_order')).scalar() == 0


def test_order_ownership_and_history_survive_backup_restore(client, headers):
    a, payload = _create_order(client, headers, 'BACKUP-A', department='甲部')
    b, _ = _create_order(client, headers, 'BACKUP-B', project_code=payload['project_code'], department='甲部')
    assert transfer(client, headers, a, account_manager='已交接负责人', department='乙部').status_code == 200
    backup = client.post('/api/backups', headers=headers)
    assert backup.status_code == 200, backup.text
    assert transfer(client, headers, a, account_manager='恢复前临时负责人').status_code == 200
    response = client.post(f"/api/backups/{backup.json()['id']}/restore", headers=headers)
    assert response.status_code == 200, response.text
    restored = client.get(f'/api/history/lines/{a}', headers=headers).json()
    assert restored['current']['account_manager'] == '已交接负责人'
    assert restored['current']['department'] == '乙部'
    assert '恢复前临时负责人' not in str(restored['managers'])
    assert client.get(f'/api/history/lines/{b}', headers=headers).json()['current']['department'] == '甲部'
    from fastapi import HTTPException
    from app.department_service import validate_restore_departments
    with db() as conn, pytest.raises(HTTPException) as error:
        validate_restore_departments(conn, {'sales_order':[{'ownership_overridden':1,'department':'未登记的交接部门'}]})
    assert error.value.status_code == 422


def test_order_transfer_rejects_stale_editor_without_partial_write(client, headers):
    from fastapi.testclient import TestClient
    from app.main import app
    from test_edit_conflicts import edit_headers
    a, _ = _create_order(client, headers, 'STALE-A', department='甲部')
    raw = TestClient(app, raise_server_exceptions=False)
    stale = edit_headers(raw, headers, a)
    current = raw.get(f'/api/history/lines/{a}', headers=headers).json()['current']
    values = {k:current[k] for k in ('account_manager','department','branch_company','team_level3_name')}
    assert transfer(client, headers, a, account_manager='已保存负责人').status_code == 200
    result = raw.post('/api/history/transfer-order', headers=stale,
                      json={'order_line_id':a, 'expected':values, **values, 'account_manager':'过期修改', 'reason':'过期测试'})
    assert result.status_code == 409, result.text
    assert result.json()['detail']['code'] == 'EDIT_CONFLICT'
    assert raw.get(f'/api/history/lines/{a}', headers=headers).json()['current']['account_manager'] == '已保存负责人'
