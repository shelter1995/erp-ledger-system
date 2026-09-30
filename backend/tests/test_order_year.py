from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import text
from app.db import db
from test_accounts_v2 import admin_v2
from test_permission_routes_v2 import seed_mixed, actor


def test_year_scope_lists_aggregates_options_exports_and_access(client, admin_v2):
    with db() as conn:
        a,b,ids=seed_mixed(conn)
        conn.execute(text("UPDATE order_line SET line_order_date='2024-12-31' WHERE id=:id"), {'id':ids[0]})
        conn.execute(text("UPDATE order_line SET line_order_date='2025-01-01' WHERE id=:id"), {'id':ids[1]})
        conn.execute(text("UPDATE order_line SET line_order_date=NULL WHERE id=:id"), {'id':ids[2]})
        conn.execute(text("UPDATE sales_order SET order_date='2026-06-01'"))
        conn.execute(text("INSERT INTO sales_receipt(order_line_id,receipt_date,receipt_amount) VALUES(:id,'2026-06-01',10)"), {'id':ids[0]})
        _,headers=actor(conn,'year_reader',['ledger_view','order_view','sale_view','purchase_view','dashboard_view'],[a])
    def get(path, **params):
        r=client.get(path,params=params,headers=admin_v2)
        assert r.status_code==200,r.text
        return r.json()
    for path in ['/api/orders','/api/sales','/api/purchases']:
        assert get(path,order_year=2024)['total']==1
        assert get(path,order_year=2025)['total']==1
        assert get(path,order_year=2026)['total']==0
        assert get(path)['total']==3
        assert get(path,order_year=2024,limit=1,offset=1)['items']==[]
        assert client.get(path,params={'order_year':1899},headers=admin_v2).status_code==422
    for path in ['/api/ledgers/summary','/api/dashboard/data']:
        assert get(path,order_year=2024)['metrics']['totalOrderAmount']=='100.00'
        assert get(path,order_year=2026)['metrics']['orderCount']==0
        assert get(path)['metrics']['totalOrderAmount']=='1700.00'
    for path in ['/api/sales/order-options','/api/purchases/order-options']:
        assert len(get(path,order_year=2024)['items'])==1
        assert len(get(path)['items'])==3
    detail=get('/api/ledgers/project-detail',project_code='AUTH-MIX',order_year=2024)
    assert detail['order_value']=='100.00'
    assert len(detail['orders'][0]['lines'])==1
    assert get('/api/sales',order_year=2024,receipt_start_date='2026-01-01')['total']==1
    assert get('/api/sales',order_year=2026,receipt_start_date='2026-01-01')['total']==0
    assert get('/api/data/order-years')['years']==[2025,2024]
    assert client.get('/api/data/order-years',headers=headers).json()['years']==[2024]
    assert client.get('/api/ledgers/summary',params={'order_year':2025},headers=headers).json()['items']==[]
    for path in ['/api/ledgers/export-summary','/api/orders/export','/api/history/export']:
        r=client.get(path,params={'order_year':2024},headers=admin_v2)
        assert r.status_code==200,r.text
        book=load_workbook(BytesIO(r.content),data_only=True)
        sheet=book['明细'] if path=='/api/history/export' else book.active
        assert sheet.max_row==(3 if path=='/api/orders/export' else 2)


def test_native_date_and_imported_date_have_same_year_semantics(client, admin_v2):
    with db() as conn:
        _,_,ids=seed_mixed(conn)
        conn.execute(text("UPDATE sales_order SET order_date='2026-01-01'"))
        conn.execute(text("UPDATE order_line SET source_preserved=0,line_order_date='2024-01-01' WHERE id=:id"),{'id':ids[0]})
        conn.execute(text("UPDATE order_line SET line_order_date='2025-12-31' WHERE id=:id"),{'id':ids[1]})
        conn.execute(text("UPDATE order_line SET line_order_date=NULL WHERE id=:id"),{'id':ids[2]})
        conn.execute(text("INSERT INTO purchase_payment(order_line_id,payment_date,payment_amount) VALUES(:id,'2026-03-01',10)"),{'id':ids[1]})
    r=client.get('/api/orders',params={'order_year':2026},headers=admin_v2)
    assert r.status_code==200,r.text
    assert [item['order_line_id'] for item in r.json()['items']]==[ids[0]]
    assert client.get('/api/data/order-years',headers=admin_v2).json()['years']==[2026,2025]
    for year,total in [(2025,1),(2026,0)]:
        r=client.get('/api/purchases',params={'order_year':year,'payment_start_date':'2026-01-01'},headers=admin_v2)
        assert r.status_code==200,r.text
        assert r.json()['total']==total
