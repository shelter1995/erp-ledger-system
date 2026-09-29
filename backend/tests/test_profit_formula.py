from decimal import Decimal
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import db
from test_source_import import source_file_with_identities
from test_template_0916 import workbook_bytes
from test_formula_linkage import batch_item


def import_profit_rows(client, headers, rows):
    workbook = load_workbook(BytesIO(source_file_with_identities([
        ('P-PROFIT', 'SO-PROFIT', f'设备-{i}') for i in range(len(rows))
    ])))
    for no, values in enumerate(rows, 3):
        for column, value in zip(('W', 'AC', 'BL', 'BM', 'BN'), values):
            workbook.active[f'{column}{no}'] = value
    response = client.post('/api/orders/source-import?filename=profit.xlsx&preview=false',
                           content=workbook_bytes(workbook), headers=headers)
    assert response.status_code == 200, response.text
    return client.get('/api/orders', headers=headers).json()['items']


@pytest.mark.parametrize('old_profit', [None, 999999])
def test_profit_uses_four_operands_ignoring_bn_and_export_keeps_tax_refund(client, headers, old_profit):
    rows = import_profit_rows(client, headers, [(100, 40, 8, 3, old_profit)])
    assert Decimal(rows[0]['gross_profit']) == Decimal('55')
    assert client.get('/api/dashboard/summary', headers=headers).json()['grossProfit'] == '55.00'
    detail = client.get('/api/sales/by-order', params={'project_id':'P-PROFIT','order_id':'SO-PROFIT'}, headers=headers)
    assert detail.status_code == 200, detail.text
    assert Decimal(detail.json()['summary']['gross_profit']) == Decimal('55')
    exported = client.get('/api/orders/export', headers=headers)
    assert exported.status_code == 200, exported.text
    book = load_workbook(BytesIO(exported.content), data_only=True)
    assert [Decimal(str(book.active[f'{col}3'].value)) for col in ('BL','BM','BN')] == [Decimal(8), Decimal(3), Decimal(55)]


def test_profit_aggregates_unrounded_source_inputs(client, headers):
    rows = import_profit_rows(client, headers, [
        (Decimal('100.004'), 40, Decimal('8.001'), 3, None),
        (Decimal('100.004'), 40, Decimal('8.001'), 3, None),
    ])
    assert all(Decimal(row['gross_profit']) == Decimal('55.003') for row in rows)
    assert client.get('/api/dashboard/summary', headers=headers).json()['grossProfit'] == '110.01'
    assert sum(Decimal(row['order_value']) for row in rows) == Decimal('200.00')


def test_profit_rounds_only_final_total_half_up(client, headers):
    import_profit_rows(client, headers, [(Decimal('100.005'), 40, 8, 3, None)])
    assert client.get('/api/dashboard/summary', headers=headers).json()['grossProfit'] == '55.01'


def test_source_profit_changes_after_sales_and_purchase_edits(client, headers):
    rows = import_profit_rows(client, headers, [(Decimal('100.004'), 40, 8, 3, 9999)])
    line = rows[0]['order_line_id']
    item = batch_item(client, headers, 'orders', line)
    item.update(quantity='1', sales_tax_rate='0', net_unit_price='200', unit_price='200')
    response = client.put('/api/orders/batch-basic', json={'items':[item]}, headers=headers)
    assert response.status_code == 200, response.text
    with db() as conn:
        row = conn.execute(text('SELECT order_value,purchase_amount,gross_profit FROM v_order_line_finance')).mappings().one()
        assert row['order_value'] == 200
        assert row['gross_profit'] == row['order_value'] - row['purchase_amount'] - 8 + 3
    item = batch_item(client, headers, 'purchases', line)
    item.update(purchase_tax_rate='0', purchase_unit_price_no_tax='500', purchase_unit_price='500')
    response = client.put('/api/purchases/batch', json={'items':[item]}, headers=headers)
    assert response.status_code == 200, response.text
    assert Decimal(client.get('/api/orders', headers=headers).json()['items'][0]['gross_profit']) == Decimal('-305')


def test_profit_upgrade_uses_archive_and_current_edited_amounts_once(client, headers):
    from app.profit_calculations import backfill_profit_inputs
    rows = import_profit_rows(client, headers, [(100, 40, 8, 3, None)])
    with db() as conn:
        conn.execute(text('''UPDATE order_line SET order_value=200, source_order_value_precise=NULL,
            source_purchase_amount_precise=NULL,profit_tax_amount=NULL,profit_tax_refund=NULL,
            profit_inputs_initialized=0'''))
        assert backfill_profit_inputs(conn) == 1
        assert conn.execute(text('SELECT gross_profit FROM v_order_line_finance')).scalar_one() == Decimal('155')
        conn.execute(text('UPDATE order_line SET profit_tax_refund=7'))
        assert backfill_profit_inputs(conn) == 0
        assert conn.execute(text('SELECT gross_profit FROM v_order_line_finance')).scalar_one() == Decimal('159')


def test_profit_bad_tax_blocks_source_import_without_partial_write(client, headers):
    book = load_workbook(BytesIO(source_file_with_identities([('P-BAD','SO-BAD','设备')])))
    book.active['BL3'] = '#VALUE!'
    response = client.post('/api/orders/source-import?filename=bad-tax.xlsx&preview=false', content=workbook_bytes(book), headers=headers)
    assert response.status_code == 422, response.text
    assert '税金' in response.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
