from io import BytesIO
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.db import db
from app.source_import import source_order_chain
from test_template_0916 import current_import_file, workbook_bytes


def test_source_history_ignores_repeated_slashes_only():
    assert source_order_chain('OLD//MIDDLE／／NEW', 3) == ['OLD', 'MIDDLE', 'NEW']
    with pytest.raises(ValueError):
        source_order_chain('/OLD/NEW', 3)


@pytest.mark.parametrize('amount,expected', [('60/40.01', ['60', '40.01']), (100.01, ['50', '50.01'])])
def test_source_delivery_phases_preserve_totals_and_platform(client, headers, amount, expected):
    wb = load_workbook(BytesIO(current_import_file()))
    ws = wb.active
    ws['M3'] = 'OLD//NEW'
    ws['AD3'] = '2025/1/2，2025/2/3'
    ws['AE3'] = 2
    ws['AF3'] = 80
    ws['AG3'] = amount
    ws['L3'] = '平台甲'
    ws['J3'] = '历史客户甲'
    ws['K3'] = '历史用户甲'
    ws['AR3'] = '2025/1/2，2025/2/3'
    ws['AS3'] = 'shared-invoice-reference'
    ws['AT3'] = 100.01
    ws.append([c.value for c in ws[3]])
    ws['AD4'] = '2025/3/4'
    ws['AG4'] = 30
    ws['L4'] = '平台乙'
    ws['J4'] = '历史客户乙'
    ws['K4'] = '历史用户乙'
    content = workbook_bytes(wb)
    preview = client.post('/api/orders/source-import?filename=phases.xlsx', content=content, headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()['summary']['delivery_amount'] == '130.01'
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
    result = client.post('/api/orders/source-import?filename=phases.xlsx&preview=false', content=content, headers=headers)
    assert result.status_code == 200, result.text
    with db() as conn:
        lines = conn.execute(text('SELECT * FROM v_order_line_finance ORDER BY order_line_id')).mappings().all()
        assert len(lines) == 2
        assert [r['regional_platform'] for r in lines] == ['平台甲','平台乙']
        assert [r['customer_unit_name'] for r in lines] == ['历史客户甲','历史客户乙']
        assert [r['end_user_name'] for r in lines] == ['历史用户甲','历史用户乙']
        first_id = lines[0]['order_line_id']
        phases = conn.execute(text('SELECT delivery_date,delivery_value FROM delivery_record WHERE order_line_id=:id ORDER BY id'), {'id':first_id}).mappings().all()
        assert [str(p['delivery_date']) for p in phases] == ['2025-01-02','2025-02-03']
        assert [p['delivery_value'] for p in phases] == [Decimal(v) for v in expected]
        assert lines[0]['delivery_value'] == Decimal('100.01')
        assert lines[0]['delivery_quantity'] == Decimal('2')
        invoices = conn.execute(text('SELECT invoice_amount,invoice_no FROM purchase_invoice WHERE order_line_id=:id ORDER BY phase_no'), {'id':first_id}).mappings().all()
        assert [r['invoice_amount'] for r in invoices] == [Decimal('50'),Decimal('50.01')]
        assert {r['invoice_no'] for r in invoices} == {'shared-invoice-reference'}
    listing = client.get('/api/orders', headers=headers).json()
    assert listing['total'] == len(listing['items']) == 2
    assert {r['regional_platform'] for r in listing['items']} == {'平台甲','平台乙'}
    detail = client.get(f'/api/sales/{first_id}', headers=headers)
    assert detail.status_code == 200
    assert len(detail.json()['deliveries']) == 2
    exported = client.get('/api/orders/export', headers=headers)
    assert exported.status_code == 200
    book = load_workbook(BytesIO(exported.content), data_only=True)
    exported_rows = list(book.active.iter_rows(min_row=3, values_only=True))
    assert len(exported_rows) == 2
    assert {row[9] for row in exported_rows} == {'历史客户甲','历史客户乙'}
    assert {row[10] for row in exported_rows} == {'历史用户甲','历史用户乙'}
    assert {row[11] for row in exported_rows} == {'平台甲', '平台乙'}
    first = next(row for row in exported_rows if row[11] == '平台甲')
    assert first[29] == '2025-01-02;2025-02-03'
    assert [Decimal(v) for v in first[32].split('/')] == [Decimal(v) for v in expected]
    book.close()
    from app.routers.orders import _upsert_line_record
    from app.routers.purchases import _upsert_order_line_record
    from fastapi import HTTPException
    with db() as conn:
        # A full-form save containing unchanged aggregate fields must preserve
        # the original phases; editing aggregate delivery values must reject.
        _upsert_line_record(conn, 'delivery_record', first_id, {'delivery_value':Decimal('100.01')})
        with pytest.raises(HTTPException) as exc:
            _upsert_line_record(conn, 'delivery_record', first_id, {'delivery_value':Decimal('200')})
        assert exc.value.status_code == 409
        assert conn.execute(text('SELECT SUM(delivery_value) FROM delivery_record WHERE order_line_id=:id'), {'id':first_id}).scalar_one() == Decimal('100.01')
        _upsert_order_line_record(conn, 'delivery_record', first_id, {'delivery_value':Decimal('100.01')})
        assert conn.execute(text('SELECT SUM(delivery_value) FROM delivery_record WHERE order_line_id=:id'), {'id':first_id}).scalar_one() == Decimal('100.01')
        with pytest.raises(HTTPException) as exc:
            _upsert_order_line_record(conn, 'delivery_record', first_id, {'delivery_value':Decimal('200')})
        assert exc.value.status_code == 409
    from test_formula_linkage import batch_item
    basic = batch_item(client, headers, 'orders', first_id)
    basic['regional_platform'] = '平台丙'
    basic['customer_unit_name'] = '历史客户丙'
    basic['user_name'] = '历史用户丙'
    response = client.put('/api/orders/batch-basic', json={'items':[basic]}, headers=headers)
    assert response.status_code == 200, response.text
    assert {r['regional_platform'] for r in client.get('/api/orders', headers=headers).json()['items']} == {'平台丙', '平台乙'}
    items = client.get('/api/orders', headers=headers).json()['items']
    assert {r['customer_unit_name'] for r in items} == {'历史客户丙','历史客户乙'}
    with db() as conn:
        assert set(conn.execute(text('SELECT end_user_name FROM v_order_line_finance')).scalars()) == {'历史用户丙','历史用户乙'}
    purchase = batch_item(client, headers, 'purchases', first_id)
    purchase['supplier_name'] = 'supplier-updated'
    response = client.put('/api/purchases/batch', json={'items':[purchase]}, headers=headers)
    assert response.status_code == 200, response.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM delivery_record WHERE order_line_id=:id'), {'id':first_id}).scalar_one() == 2
        assert conn.execute(text('SELECT SUM(delivery_value) FROM delivery_record WHERE order_line_id=:id'), {'id':first_id}).scalar_one() == Decimal('100.01')


def test_source_delivery_mismatched_sequences_reject_without_writes(client, headers):
    wb = load_workbook(BytesIO(current_import_file()))
    wb.active['AD3'] = '2025/1/2，2025/2/3'
    wb.active['AG3'] = '10/20/30'
    result = client.post('/api/orders/source-import?filename=mismatch.xlsx', content=workbook_bytes(wb), headers=headers)
    assert result.status_code == 422
    assert 'AG3' in result.json()['detail']


def test_source_preserves_long_purchase_contract_reference(client, headers):
    wb = load_workbook(BytesIO(current_import_file()))
    reference = 'CONTRACT-A-' * 10
    wb.active['AM3'] = reference
    result = client.post('/api/orders/source-import?filename=long-contract.xlsx&preview=false', content=workbook_bytes(wb), headers=headers)
    assert result.status_code == 200, result.text
    with db() as conn:
        assert conn.execute(text('SELECT purchase_contract_no FROM purchase_contract')).scalar_one() == reference


@pytest.mark.parametrize('paths,expected_orders', [(['OLD-A/NEW','OLD-B/NEW'],1), (['OLD/NEW-A','OLD/NEW-B'],2), (['OLD/MID','OLD/MID/NEW'],2)])
def test_source_order_mergers_and_splits_preserve_each_line_path(client, headers, paths, expected_orders):
    from test_source_import import source_file_with_identities
    content=source_file_with_identities([('P-MOVEMENT', path, f'goods-{i}') for i,path in enumerate(paths)])
    preview=client.post('/api/orders/source-import?filename=movement.xlsx',content=content,headers=headers)
    assert preview.status_code==200,preview.text
    assert preview.json()['summary']['order_count']==expected_orders
    result=client.post('/api/orders/source-import?filename=movement.xlsx&preview=false',content=content,headers=headers)
    assert result.status_code==200,result.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one()==expected_orders
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one()==2
    items=client.get('/api/orders',headers=headers).json()['items']
    assert {r['goods_name']:r['order_number_path'] for r in items}=={f'goods-{i}':path.split('/') for i,path in enumerate(paths)}
    assert {r['order_no'] for r in items}=={p.split('/')[-1] for p in paths}
    for path in paths:
        found=client.get('/api/orders',params={'order_id':path.split('/')[0]},headers=headers)
        assert found.status_code==200 and found.json()['total']==2
    export=client.get('/api/orders/export',headers=headers)
    book=load_workbook(BytesIO(export.content),data_only=True)
    assert {r[14]:r[12] for r in book.active.iter_rows(min_row=3,values_only=True)}=={f'goods-{i}':path for i,path in enumerate(paths)}
    book.close()


def test_source_split_does_not_move_existing_order_lines(client,headers):
    from test_source_import import source_file_with_identities
    for filename,paths in [('existing.xlsx',['OLD']),('split.xlsx',['OLD/NEW-A','OLD/NEW-B'])]:
        content=source_file_with_identities([('P-SPLIT',path,f'{filename}-{i}') for i,path in enumerate(paths)])
        result=client.post(f'/api/orders/source-import?filename={filename}&preview=false',content=content,headers=headers)
        assert result.status_code==200,result.text
    items=client.get('/api/orders',params={'order_id':'OLD'},headers=headers).json()['items']
    assert len(items)==3
    assert {r['order_no'] for r in items}=={'OLD','NEW-A','NEW-B'}
    original=next(r for r in items if r['goods_name']=='existing.xlsx-0')
    assert original['order_no']=='OLD' and original['order_number_path']==['OLD']


def test_source_long_invoice_references_are_not_truncated(client,headers):
    wb=load_workbook(BytesIO(current_import_file()))
    reference='/'.join(f'INV-{i:016}' for i in range(20))
    wb.active['AS3']=reference
    wb.active['BW3']=reference
    wb.active.cell(3,73,reference)
    wb.active.cell(3,80,reference)
    result=client.post('/api/orders/source-import?filename=invoices.xlsx&preview=false',content=workbook_bytes(wb),headers=headers)
    assert result.status_code==200,result.text
    with db() as conn:
        for table in ('purchase_invoice','sales_invoice'):
            assert conn.execute(text(f'SELECT invoice_no FROM {table}')).scalar_one()==reference
        assert conn.execute(text('SELECT invoice_doc_no FROM sales_invoice')).scalar_one()==reference
        assert conn.execute(text('SELECT payment_notice_no FROM sales_receipt WHERE phase_no=1')).scalar_one()==reference
