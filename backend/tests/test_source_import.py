from io import BytesIO
from decimal import Decimal, ROUND_HALF_UP
import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import event, text
from app.db import db, engine
from test_financial_integration import _excel_import_file
from test_template_0916 import current_import_file, workbook_bytes
from test_import_permissions import _create_user, _login


def source_file():
    wb = load_workbook(BytesIO(current_import_file()))
    ws = wb.active
    ws['E3'] = '甲'
    ws['M3'] = 'A/B'
    ws['BV3'] = '2026/8/1，2026/8/2'
    ws['BX3'] = 200
    ws['BW3'] = '票1/票2'
    ws['AU3'] = None
    ws['AW3'] = 50
    values = [c.value for c in ws[3]]
    ws.append(values)
    ws['E4'] = '乙'
    ws['C4'] = '第二部门'
    ws['I4'] = '第二团队'
    return workbook_bytes(wb)


def source_file_with_identities(rows):
    """Build source rows without using any business workbook data."""
    wb = load_workbook(BytesIO(current_import_file()))
    ws = wb.active
    template = [cell.value for cell in ws[3]]
    for index, (project_code, order_no, goods_name) in enumerate(rows, 3):
        if index > 3:
            ws.append(template)
        ws.cell(index, 2).value = project_code
        ws.cell(index, 13).value = order_no
        ws.cell(index, 15).value = goods_name
        ws.cell(index, 16).value = f'规格-{index}'
    return workbook_bytes(wb)


@pytest.mark.parametrize('amount', [0, 0.25, 123.45, 12345678])
def test_source_import_recovers_numbers_with_date_formats(client, headers, amount):
    wb = load_workbook(BytesIO(current_import_file()))
    ws = wb.active
    ws['AR3'] = '2025-01-02'
    ws['AT3'] = amount
    ws['AS3'] = 12345678
    ws['BW3'] = 87654321
    for coordinate in ('AT3', 'AS3', 'BW3'):
        ws[coordinate].number_format = 'yyyy-mm-dd'
    content = workbook_bytes(wb)
    preview = client.post('/api/orders/source-import?filename=date-format.xlsx', content=content, headers=headers)
    assert preview.status_code == 200, preview.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
    result = client.post('/api/orders/source-import?filename=date-format.xlsx&preview=false', content=content, headers=headers)
    assert result.status_code == 200, result.text
    assert any('日期格式' in item['message'] for item in result.json()['warnings'])
    with db() as conn:
        row = conn.execute(text('SELECT invoice_amount,invoice_no,received_invoice_date FROM purchase_invoice')).mappings().one()
        assert row['invoice_amount'] == Decimal(str(amount))
        assert row['invoice_no'] == '12345678'
        assert row['received_invoice_date'].isoformat() == '2025-01-02'
        assert conn.execute(text('SELECT invoice_no FROM sales_invoice')).scalar_one() == '87654321'


def test_source_import_invalid_finance_identifies_cell_without_echoing_value(client, headers):
    wb = load_workbook(BytesIO(current_import_file()))
    wb.active['AT3'] = 'private-invalid-amount'
    response = client.post('/api/orders/source-import?filename=invalid.xlsx', content=workbook_bytes(wb), headers=headers)
    assert response.status_code == 422
    detail = response.json()['detail']
    assert 'AT3' in detail and '收票金额' in detail
    assert 'private-invalid-amount' not in detail


def test_source_preview_batches_post_import_line_work(client, headers):
    content = source_file_with_identities([
        (f'P-BATCH-{index}', f'SO-BATCH-{index}', f'设备-{index}')
        for index in range(6)
    ])
    def captured_post(url):
        statements = []

        def observe(_conn, _cursor, statement, _parameters, _context, _executemany):
            statements.append(' '.join(statement.lower().split()))

        event.listen(engine, 'before_cursor_execute', observe)
        try:
            response = client.post(url, content=content, headers=headers)
        finally:
            event.remove(engine, 'before_cursor_execute', observe)
        return response, statements

    def assert_no_per_line_finalization(statements):
        assert not [
            statement for statement in statements
            if 'from v_order_line_finance where order_line_id=' in statement
        ]
        assert not [
            statement for statement in statements
            if 'select so.project_id from order_line ol join sales_order so' in statement
            and 'where ol.id=' in statement
        ]
        assert not [
            statement for statement in statements
            if 'sales_order_number_history' in statement
            and 'where sales_order_id = ' in statement
        ]
        assert not [
            statement for statement in statements
            if 'select distinct p.id, p.project_code from sales_order so' in statement
            and 'where so.order_no=' in statement
        ]
        assert not [
            statement for statement in statements
            if 'select id, deleted_at from project where project_code=' in statement
        ]
        forbidden_writes = (
            'update order_line set source_preserved=1',
            'update ledger_raw_row r join order_line ol',
            'update sales_invoice set invoice_doc_no=',
            'update purchase_payment set due_payment_date=',
            'delete from purchase_payment where order_line_id=',
            'update order_line set sales_unit_price_no_tax=',
            'update purchase_info set purchase_unit_price_no_tax=',
            'update delivery_record set delivery_revenue_no_tax=',
        )
        assert not [
            statement for statement in statements
            if any(fragment in statement for fragment in forbidden_writes)
            and 'case ' not in statement
        ]
        phase_tables = (
            'purchase_invoice',
            'warehouse_entry',
            'finance_payment_entry',
            'purchase_payment',
            'sales_invoice',
            'sales_receipt',
        )
        for table in phase_tables:
            assert len([
                statement for statement in statements
                if statement.startswith(f'insert into {table} ')
            ]) <= 1

    preview, preview_statements = captured_post(
        '/api/orders/source-import?filename=batch-preview.xlsx'
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()['success_rows'] == 6
    assert_no_per_line_finalization(preview_statements)

    imported, import_statements = captured_post(
        '/api/orders/source-import?filename=batch-preview.xlsx&preview=false'
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()['success_rows'] == 6
    assert_no_per_line_finalization(import_statements)

    with db() as conn:
        invoices = conn.execute(text('''SELECT si.pending_invoice_amount,
            si.delivered_not_invoiced_amount,v.order_value,v.delivery_value,
            v.sales_invoice_amount FROM sales_invoice si
            JOIN v_order_line_finance v ON v.order_line_id=si.order_line_id''')).mappings().all()
        receipts = conn.execute(text('''SELECT sr.receipt_amount,sr.receipt_ratio,
            v.sales_invoice_amount FROM sales_receipt sr
            JOIN v_order_line_finance v ON v.order_line_id=sr.order_line_id''')).mappings().all()
        orders = conn.execute(text('''SELECT so.close_status,v.accounts_receivable
            FROM sales_order so JOIN order_line ol ON ol.sales_order_id=so.id
            JOIN v_order_line_finance v ON v.order_line_id=ol.id
            WHERE so.order_no LIKE 'SO-BATCH-%' ''')).mappings().all()
    assert invoices
    assert all(
        row['pending_invoice_amount'] == row['order_value'] - row['sales_invoice_amount']
        and row['delivered_not_invoiced_amount'] == row['delivery_value'] - row['sales_invoice_amount']
        for row in invoices
    )
    assert receipts
    assert all(
        row['receipt_ratio'] == (
            row['receipt_amount'] / row['sales_invoice_amount'] * 100
        ).quantize(Decimal('.000001'), rounding=ROUND_HALF_UP)
        for row in receipts if row['sales_invoice_amount']
    )
    assert len(orders) == 6
    assert all(
        row['close_status'] == (None if row['accounts_receivable'] else '关闭')
        for row in orders
    )


def transitional_91_source_file():
    """Build the observed 2024 layout without using any business workbook data."""
    wb = load_workbook(BytesIO(_excel_import_file()))
    ws = wb.active
    ws.delete_cols(25)
    ws.delete_cols(87)
    ws.insert_cols(87, 2)
    ws.cell(2, 87, '交付应收款')
    ws.cell(2, 88, '开票应收款')
    ws.cell(3, 87, Decimal('999999'))
    ws.cell(3, 88, Decimal('888888'))
    assert ws.max_column == 91
    return workbook_bytes(wb)


def transitional_92_source_file():
    """Build the observed post-Y-column layout without using business data."""
    wb = load_workbook(BytesIO(_excel_import_file()))
    ws = wb.active
    ws.delete_cols(88)
    ws.insert_cols(88, 2)
    ws.cell(2, 88, '交付应收款')
    ws.cell(2, 89, '开票应收款')
    ws.cell(3, 88, Decimal('999999'))
    ws.cell(3, 89, Decimal('888888'))
    assert ws.max_column == 92
    return workbook_bytes(wb)


def test_transitional_91_source_import_maps_columns_without_financial_shift(client, headers):
    content = transitional_91_source_file()
    preview = client.post('/api/orders/source-import?filename=2024.xlsx', content=content, headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()['layout'] == '过渡版台账'
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0

    imported = client.post(
        '/api/orders/source-import?filename=2024.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    with db() as conn:
        row = conn.execute(text(
            'SELECT v.order_value,v.purchase_unit_price_no_tax,v.purchase_unit_price,'
            'v.purchase_amount,p.purchase_tax_rate,p.labor_cost,p.other_cost '
            'FROM v_order_line_finance v JOIN purchase_info p ON p.order_line_id=v.order_line_id'
        )).mappings().one()
    assert row['order_value'] == Decimal('226.00')
    assert row['purchase_tax_rate'] is None
    assert row['purchase_unit_price_no_tax'] == Decimal('70.000000')
    assert row['purchase_unit_price'] == Decimal('79.100000')
    assert row['purchase_amount'] == Decimal('158.20')
    assert row['labor_cost'] == Decimal('12.34')
    assert row['other_cost'] == Decimal('5.67')


def test_transitional_92_source_import_preserves_purchase_tax_and_financial_columns(client, headers):
    content = transitional_92_source_file()
    preview = client.post('/api/orders/source-import?filename=after-y.xlsx', content=content, headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()['layout'] == '过渡版台账（含采购税率）'

    imported = client.post(
        '/api/orders/source-import?filename=after-y.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    with db() as conn:
        row = conn.execute(text(
            'SELECT v.order_value,v.purchase_unit_price_no_tax,v.purchase_unit_price,'
            'v.purchase_amount,p.purchase_tax_rate,p.labor_cost,p.other_cost '
            'FROM v_order_line_finance v JOIN purchase_info p ON p.order_line_id=v.order_line_id'
        )).mappings().one()
    assert row['order_value'] == Decimal('226.00')
    assert row['purchase_tax_rate'] == Decimal('13.000000')
    assert row['purchase_unit_price_no_tax'] == Decimal('70.000000')
    assert row['purchase_unit_price'] == Decimal('79.100000')
    assert row['purchase_amount'] == Decimal('158.20')
    assert row['labor_cost'] == Decimal('12.34')
    assert row['other_cost'] == Decimal('5.67')


def test_source_import_ignores_bn_result_in_dashboard_and_sales_detail(client, headers):
    workbook = load_workbook(BytesIO(transitional_92_source_file()))
    workbook.active['BN3'] = Decimal('12.34')
    workbook.active['BL3'] = Decimal('7.80')
    workbook.active['BM3'] = Decimal('3.00')
    content = workbook_bytes(workbook)

    imported = client.post(
        '/api/orders/source-import?filename=bn-gross-profit.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text

    with db() as conn:
        row = conn.execute(text(
            'SELECT project_code, order_no, source_gross_profit, gross_profit '
            'FROM order_line ol JOIN v_order_line_finance v ON v.order_line_id = ol.id'
        )).mappings().one()
    assert row['source_gross_profit'] == Decimal('12.34')
    assert row['gross_profit'] == Decimal('63.00')  # 226 - 158.20 - 7.80 + 3

    dashboard = client.get('/api/dashboard/summary', headers=headers)
    assert dashboard.status_code == 200, dashboard.text
    assert dashboard.json()['grossProfit'] == '63.00'

    detail = client.get(
        '/api/sales/by-order',
        params={'project_id': row['project_code'], 'order_id': row['order_no']},
        headers=headers,
    )
    assert detail.status_code == 200, detail.text
    assert Decimal(detail.json()['summary']['gross_profit']) == Decimal('63.00')


def test_source_rows_preview_commit_and_duplicate_file(client, headers):
    content = source_file()
    preview = client.post('/api/orders/source-import?filename=real.xlsx',content=content,headers=headers)
    assert preview.status_code == 200, preview.text
    assert preview.json()['success_rows'] == 2
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
    imported = client.post('/api/orders/source-import?filename=real.xlsx&preview=false',content=content,headers=headers)
    assert imported.status_code == 200, imported.text
    with db() as conn:
        lines = conn.execute(text('SELECT * FROM v_order_line_finance ORDER BY order_line_id')).mappings().all()
        assert [r['account_manager'] for r in lines] == ['甲','乙']
        assert [r['order_no'] for r in lines] == ['B','B']
        assert lines[1]['department'] == '第二部门'
        assert conn.execute(text(
            'SELECT order_no FROM sales_order_number_history ORDER BY history_order'
        )).scalars().all() == ['A', 'B']
        invoices = conn.execute(text('SELECT * FROM sales_invoice')).mappings().all()
        assert len(invoices) == 2
        assert all(r['invoice_date'] is None and r['invoice_amount'] == Decimal('200') and r['invoice_date_text'] == '2026/8/1，2026/8/2' for r in invoices)
        assert conn.execute(text('SELECT COUNT(*) FROM project_manager_history')).scalar_one() == 0
    repeat = client.post('/api/orders/source-import?filename=real.xlsx&preview=false',content=content,headers=headers)
    assert repeat.status_code == 409
    # List and export must expose row ownership, never the first manager on a framework.
    listed = client.get('/api/orders',headers=headers).json()['items']
    assert {r['account_manager'] for r in listed} == {'甲','乙'}
    export = client.get('/api/orders/export',headers=headers)
    book = load_workbook(BytesIO(export.content),data_only=True)
    assert {r[4] for r in book.active.iter_rows(min_row=3,values_only=True)} == {'甲','乙'}
    assert _create_user(client,headers,'source_scope',['order_edit'],department_scope=['第二部门'],department_all=False).status_code == 200
    scoped = _login(client,'source_scope')
    visible = client.get('/api/orders',headers=scoped).json()
    assert visible['total'] == 1 and visible['items'][0]['account_manager'] == '乙'
    ledger = client.get('/api/ledgers',headers=scoped)
    assert ledger.status_code == 200,ledger.text
    assert len(ledger.json()['items']) == 1
    assert Decimal(ledger.json()['items'][0]['order_amount']) == lines[1]['order_value']
    assert ledger.json()['items'][0]['account_manager'] == '乙'
    assert client.get(f"/api/history/lines/{lines[0]['order_line_id']}",headers=scoped).status_code == 404


def test_source_import_persists_long_order_number_chain_as_history(client, headers):
    chain = [
        'XSDD2024102301233',
        'XSDD2024102301232',
        'XSDD2024103000181',
        'XSDD2025022100984',
    ]
    raw_order_no = '/'.join(chain)
    assert len(raw_order_no) > 64
    content = source_file_with_identities([(None, raw_order_no, '文香项目设备')])

    preview = client.post(
        '/api/orders/source-import?filename=2024.xlsx', content=content, headers=headers
    )
    assert preview.status_code == 200, preview.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM import_batch')).scalar_one() == 0
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 0
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order_number_history')).scalar_one() == 0

    imported = client.post(
        '/api/orders/source-import?filename=2024.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    with db() as conn:
        order = conn.execute(text('SELECT id,order_no FROM sales_order')).mappings().one()
        history = conn.execute(text(
            'SELECT order_no FROM sales_order_number_history '
            'WHERE sales_order_id=:id ORDER BY history_order'
        ), {'id': order['id']}).scalars().all()
        raw_json = conn.execute(text('SELECT raw_json FROM ledger_raw_row')).scalar_one()
    assert order['order_no'] == chain[-1]
    assert history == chain
    raw_payload = json.loads(raw_json) if isinstance(raw_json, str) else raw_json
    assert raw_payload['values'][12] == raw_order_no

    found = client.get('/api/orders', params={'order_id': chain[0]}, headers=headers)
    assert found.status_code == 200, found.text
    assert found.json()['total'] == 1
    assert found.json()['items'][0]['order_no'] == chain[-1]
    assert found.json()['items'][0]['order_number_history'] == chain


def test_source_import_rejects_single_overlong_order_number_without_500(client, headers):
    content = source_file_with_identities([('P-LONG-ORDER', 'X' * 65, '设备')])

    response = client.post(
        '/api/orders/source-import?filename=overlong.xlsx', content=content, headers=headers
    )

    assert response.status_code == 422, response.text
    assert '第 3 行' in response.text
    assert '销售订单号' in response.text
    assert '64' in response.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM import_batch')).scalar_one() == 0
        assert conn.execute(text('SELECT COUNT(*) FROM ledger_raw_row')).scalar_one() == 0


def test_blank_project_code_groups_by_order_and_preserves_original_cell(client, headers):
    content = source_file_with_identities([
        (None, 'SO-HISTORY-001', '核心交换机'),
        (None, 'SO-HISTORY-001', '配套模块'),
    ])
    preview = client.post('/api/orders/source-import?filename=history.xlsx', content=content, headers=headers)
    assert preview.status_code == 200, preview.text
    assert any('SO-HISTORY-001' in warning['message'] and '临时待补-' in warning['message']
               for warning in preview.json()['warnings'])

    imported = client.post(
        '/api/orders/source-import?filename=history.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    with db() as conn:
        projects = conn.execute(text('SELECT project_code FROM project WHERE deleted_at IS NULL')).scalars().all()
        assert len(projects) == 1
        assert projects[0].startswith('临时待补-核心交换机等项目-')
        assert len(projects[0]) <= 64
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 1
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 2
        raw_rows = conn.execute(text('SELECT raw_json FROM ledger_raw_row ORDER BY excel_row_no')).scalars().all()
    for raw in raw_rows:
        payload = json.loads(raw) if isinstance(raw, str) else raw
        assert payload['values'][1] is None


def test_blank_project_code_does_not_merge_same_first_goods_across_orders(client, headers):
    content = source_file_with_identities([
        (None, 'SO-HISTORY-A', '通用设备'),
        (None, 'SO-HISTORY-B', '通用设备'),
    ])
    imported = client.post(
        '/api/orders/source-import?filename=two-orders.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert imported.status_code == 200, imported.text
    with db() as conn:
        codes = conn.execute(text('SELECT project_code FROM project ORDER BY project_code')).scalars().all()
    assert len(codes) == 2
    assert all(code.startswith('临时待补-通用设备等项目-') for code in codes)
    assert len(set(codes)) == 2


def test_blank_project_code_rejects_order_number_owned_by_multiple_projects(client, headers):
    with db() as conn:
        for code in ('P-AMBIGUOUS-1', 'P-AMBIGUOUS-2'):
            project_id = conn.execute(
                text('INSERT INTO project (project_code) VALUES (:code)'), {'code': code}
            ).lastrowid
            conn.execute(
                text('INSERT INTO sales_order (project_id,order_no) VALUES (:project_id,:order_no)'),
                {'project_id': project_id, 'order_no': 'SO-AMBIGUOUS'},
            )
    content = source_file_with_identities([(None, 'SO-AMBIGUOUS', '设备')])
    response = client.post('/api/orders/source-import?filename=ambiguous.xlsx', content=content, headers=headers)
    assert response.status_code == 422, response.text
    assert 'SO-AMBIGUOUS' in response.text and '多个项目' in response.text


def test_later_official_code_renames_single_temporary_project(client, headers):
    historical = source_file_with_identities([(None, 'SO-UPGRADE', '历史设备')])
    first = client.post(
        '/api/orders/source-import?filename=historical.xlsx&preview=false',
        content=historical,
        headers=headers,
    )
    assert first.status_code == 200, first.text

    official = source_file_with_identities([('P-OFFICIAL', 'SO-UPGRADE', '追加设备')])
    second = client.post(
        '/api/orders/source-import?filename=official.xlsx&preview=false',
        content=official,
        headers=headers,
    )
    assert second.status_code == 200, second.text
    with db() as conn:
        assert conn.execute(text('SELECT project_code FROM project')).scalars().all() == ['P-OFFICIAL']
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 1
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 2


def test_later_official_code_merges_safe_single_order_temporary_project(client, headers):
    official = source_file_with_identities([('P-OFFICIAL', 'SO-OTHER', '正式项目设备')])
    historical = source_file_with_identities([(None, 'SO-MERGE', '历史设备')])
    assert client.post(
        '/api/orders/source-import?filename=official-base.xlsx&preview=false',
        content=official,
        headers=headers,
    ).status_code == 200
    assert client.post(
        '/api/orders/source-import?filename=historical-base.xlsx&preview=false',
        content=historical,
        headers=headers,
    ).status_code == 200

    reconciliation = source_file_with_identities([('P-OFFICIAL', 'SO-MERGE', '追加设备')])
    response = client.post(
        '/api/orders/source-import?filename=reconcile.xlsx&preview=false',
        content=reconciliation,
        headers=headers,
    )
    assert response.status_code == 200, response.text
    with db() as conn:
        assert conn.execute(text('SELECT project_code FROM project')).scalars().all() == ['P-OFFICIAL']
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 2
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 3


def test_explicit_project_codes_keep_same_order_number_in_separate_projects(client, headers):
    first = source_file_with_identities([('P-FRAMEWORK-1', 'SO-REUSED', '设备一')])
    second = source_file_with_identities([('P-FRAMEWORK-2', 'SO-REUSED', '设备二')])
    assert client.post(
        '/api/orders/source-import?filename=framework-1.xlsx&preview=false',
        content=first,
        headers=headers,
    ).status_code == 200
    response = client.post(
        '/api/orders/source-import?filename=framework-2.xlsx&preview=false',
        content=second,
        headers=headers,
    )
    assert response.status_code == 200, response.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM project')).scalar_one() == 2
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 2


def test_same_file_keeps_reused_order_number_under_explicit_projects(client, headers):
    content = source_file_with_identities([
        ('P-FRAMEWORK-1', 'SO-REUSED', '设备一'),
        ('P-FRAMEWORK-2', 'SO-REUSED', '设备二'),
    ])
    response = client.post(
        '/api/orders/source-import?filename=two-frameworks.xlsx&preview=false',
        content=content,
        headers=headers,
    )
    assert response.status_code == 200, response.text
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM project')).scalar_one() == 2
        assert conn.execute(text('SELECT COUNT(*) FROM sales_order')).scalar_one() == 2


def test_same_current_order_number_keeps_distinct_histories_across_projects(client, headers):
    content = source_file_with_identities([
        ('P-FRAMEWORK-1', 'OLD-1/SO-REUSED', '设备一'),
        ('P-FRAMEWORK-2', 'OLD-2/SO-REUSED', '设备二'),
    ])

    response = client.post(
        '/api/orders/source-import?filename=separate-histories.xlsx&preview=false',
        content=content,
        headers=headers,
    )

    assert response.status_code == 200, response.text
    with db() as conn:
        histories = conn.execute(text('''SELECT p.project_code,h.order_no,h.history_order
            FROM sales_order_number_history h
            JOIN sales_order so ON so.id=h.sales_order_id
            JOIN project p ON p.id=so.project_id
            ORDER BY p.project_code,h.history_order''')).all()
    assert histories == [
        ('P-FRAMEWORK-1', 'OLD-1', 1),
        ('P-FRAMEWORK-1', 'SO-REUSED', 2),
        ('P-FRAMEWORK-2', 'OLD-2', 1),
        ('P-FRAMEWORK-2', 'SO-REUSED', 2),
    ]


@pytest.mark.skipif(not os.environ.get('REAL_SOURCE_WORKBOOK'), reason='local acceptance workbook not configured')
def test_real_source_acceptance(client, headers):
    content = Path(os.environ['REAL_SOURCE_WORKBOOK']).read_bytes()
    response = client.post('/api/orders/source-import?filename=acceptance.xlsx&preview=false',content=content,headers=headers)
    assert response.status_code == 200, response.text
    assert response.json()['success_rows'] == 1048
    from decimal import ROUND_HALF_UP
    source = load_workbook(BytesIO(content),data_only=True).worksheets[0]
    with db() as conn:
        rows = conn.execute(text('SELECT v.*,ol.source_excel_row_no FROM v_order_line_finance v JOIN order_line ol ON ol.id=v.order_line_id')).mappings().all()
        assert len(rows) == 1048
        for row in rows:
            no = row['source_excel_row_no']
            for col,key in [(2,'project_code'),(3,'department'),(4,'branch_company'),(5,'account_manager'),(9,'team_level3_name'),(13,'order_no')]:
                assert row[key] == str(source.cell(no,col).value).strip(), (no,key)
            for col,key in [(22,'revenue_no_tax'),(23,'order_value'),(28,'cost_no_tax'),(29,'purchase_amount'),(33,'delivery_value'),(76,'sales_invoice_amount')]:
                value = source.cell(no,col).value
                expected = Decimal(str(value)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP) if value is not None else None
                assert (row[key] or 0) == (expected or 0), (no,key,row[key],expected)
            receipts = sum((Decimal(str(source.cell(no,col).value or 0)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP) for col in (81,85)),Decimal(0))
            assert row['total_received'] == receipts, no
            assert row['delivery_accounts_receivable'] == (row['delivery_value'] or 0)-receipts
            assert row['invoice_accounts_receivable'] == (row['sales_invoice_amount'] or 0)-receipts
    assert client.get('/api/orders',headers=headers).json()['total'] == 1048
    assert client.get('/api/ledgers?limit=500',headers=headers).json()['total'] == 85


def test_source_import_failure_rolls_back_every_row(client, headers):
    wb = load_workbook(BytesIO(source_file()))
    wb.active['F4'] = 'bad date'
    result = client.post('/api/orders/source-import?filename=bad.xlsx&preview=false',content=workbook_bytes(wb),headers=headers)
    assert result.status_code == 422
    with db() as conn:
        assert conn.execute(text('SELECT COUNT(*) FROM order_line')).scalar_one() == 0
        assert conn.execute(text('SELECT COUNT(*) FROM import_batch')).scalar_one() == 0
