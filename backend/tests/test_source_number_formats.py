"""Synthetic workbook format regressions; never read business spreadsheets."""
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook, load_workbook

from app.source_import import restore_source_number_formats, source_column


@pytest.mark.parametrize('layout', [91, 92, 'transitional_91', 'transitional_92'])
def test_restore_uses_sheet_relationships_and_source_layout(layout):
    wb = Workbook()
    wb.active.title = 'Instructions'
    ws = wb.create_sheet('Business data')
    columns = [46, 75, 91]
    for column, value in zip(columns, [123.123456789, 87654321, 12.34]):
        cell = ws.cell(3, source_column(layout, column), value)
        cell.number_format = 'yyyy-mm-dd'
    date_cell = ws.cell(3, source_column(layout, 44), datetime(2025, 1, 2))
    date_cell.number_format = 'yyyy-mm-dd'
    error = ws.cell(4, source_column(layout, 46), '#VALUE!')
    error.number_format = 'yyyy-mm-dd'
    original = BytesIO(); wb.save(original); wb.close()
    content = original.getvalue()
    with pytest.warns(UserWarning):
        loaded = load_workbook(BytesIO(content), data_only=True)
    ws = loaded['Business data']
    warnings = restore_source_number_formats(content, ws, layout)
    assert len(warnings) == 3
    assert ws.cell(3, source_column(layout, 46)).value == 123.123456789
    assert ws.cell(3, source_column(layout, 75)).value == '87654321'
    assert ws.cell(3, source_column(layout, 91)).value == 12.34
    assert ws.cell(3, source_column(layout, 44)).value == datetime(2025, 1, 2)
    assert ws.cell(4, source_column(layout, 46)).value == '#VALUE!'
    normalized = BytesIO(); loaded.save(normalized); loaded.close()
    reread = load_workbook(BytesIO(normalized.getvalue()), data_only=True)
    assert reread['Business data'].cell(3, source_column(layout, 46)).value == 123.123456789
    reread.close()


def test_restore_uses_cached_formula_number_without_evaluating_formula():
    wb = Workbook(); ws = wb.active
    ws['AT3'] = '=1/4'; ws['AT3'].number_format = 'yyyy-mm-dd'
    original = BytesIO(); wb.save(original); wb.close()
    patched = BytesIO()
    with ZipFile(BytesIO(original.getvalue())) as archive, ZipFile(patched, 'w', ZIP_DEFLATED) as output:
        for item in archive.infolist():
            content = archive.read(item.filename)
            if item.filename == 'xl/worksheets/sheet1.xml':
                root = ElementTree.fromstring(content)
                root.find('.//{*}c[@r="AT3"]/{*}v').text = '0.25'
                content = ElementTree.tostring(root)
            output.writestr(item, content)
    content = patched.getvalue()
    loaded = load_workbook(BytesIO(content), data_only=True)
    warnings = restore_source_number_formats(content, loaded.active, 92)
    assert len(warnings) == 1
    assert loaded.active['AT3'].value == 0.25
    loaded.close()
