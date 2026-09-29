from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text

from app.db import db
from test_source_import import source_file_with_identities
from test_template_0916 import workbook_bytes
from test_formula_linkage import batch_item


def import_dated_line(client, headers, year, amount, *, blank=False):
    book = load_workbook(BytesIO(source_file_with_identities([
        ('P-DATE', 'SO-DATE', f'设备-{year}'),
    ])))
    book.active['F3'] = None if blank else date(year, 6, 15)
    book.active['W3'] = amount
    response = client.post(
        f'/api/orders/source-import?filename={year}.xlsx&preview=false',
        content=workbook_bytes(book), headers=headers,
    )
    assert response.status_code == 200, response.text


def test_source_dates_survive_later_import_and_filter_exports(client, headers):
    import_dated_line(client, headers, 2025, 200)
    import_dated_line(client, headers, 2024, 100)
    import_dated_line(client, headers, 2026, 300, blank=True)
    for year, amount in ((2025, '200.00'), (2024, '100.00')):
        params = {'start_date': f'{year}-01-01', 'end_date': f'{year}-12-31'}
        response = client.get('/api/orders', params=params, headers=headers)
        assert response.status_code == 200, response.text
        rows = response.json()['items']
        assert len(rows) == 1
        assert rows[0]['order_date'] == f'{year}-06-15'
        assert rows[0]['order_value'] == amount
        export = client.get('/api/orders/export', params=params, headers=headers)
        assert export.status_code == 200, export.text
        book = load_workbook(BytesIO(export.content), data_only=True)
        exported = list(book.active.iter_rows(min_row=3, values_only=True))
        assert len(exported) == 1
        assert str(exported[0][5])[:10] == f'{year}-06-15'
    rows = client.get('/api/orders', headers=headers).json()['items']
    assert len(rows) == 3
    assert next(r for r in rows if r['goods_name'] == '设备-2026')['order_date'] is None


def test_source_date_upgrade_restores_archive_once_and_preserves_cleared_date(client, headers):
    from app.source_order_dates import backfill_source_order_dates
    from app.db import apply_runtime_migrations

    import_dated_line(client, headers, 2025, 200)
    import_dated_line(client, headers, 2024, 100)
    with db() as conn:
        epoch_before = conn.execute(text('SELECT data_epoch FROM business_state WHERE id=1')).scalar_one()
        conn.execute(text('UPDATE order_line SET line_order_date=NULL, line_order_date_initialized=0'))
    apply_runtime_migrations()
    with db() as conn:
        assert conn.execute(text('SELECT data_epoch FROM business_state WHERE id=1')).scalar_one() == epoch_before + 1
        rows = conn.execute(text('SELECT id,line_order_date FROM order_line ORDER BY id')).mappings().all()
        assert [r['line_order_date'] for r in rows] == [date(2025, 6, 15), date(2024, 6, 15)]
        conn.execute(text('UPDATE order_line SET line_order_date=NULL WHERE id=:id'), {'id': rows[0]['id']})
        backfill_source_order_dates(conn)
        assert conn.execute(text('SELECT line_order_date FROM order_line WHERE id=:id'), {'id': rows[0]['id']}).scalar() is None


def test_source_date_edits_do_not_change_siblings_or_shared_order_date(client, headers):
    import_dated_line(client, headers, 2025, 200)
    import_dated_line(client, headers, 2024, 100)
    rows = client.get('/api/orders', headers=headers).json()['items']
    ids = {r['goods_name']: r['order_line_id'] for r in rows}
    first = batch_item(client, headers, 'orders', ids['设备-2025'])
    second = batch_item(client, headers, 'orders', ids['设备-2024'])
    first['order_date'] = '2025-08-12'
    second['order_date'] = '2024-09-18'
    response = client.put('/api/orders/batch-basic', json={'items': [first, second]}, headers=headers)
    assert response.status_code == 200, response.text
    rows = client.get('/api/orders', headers=headers).json()['items']
    assert {r['goods_name']: r['order_date'] for r in rows} == {
        '设备-2025': '2025-08-12', '设备-2024': '2024-09-18',
    }
    with db() as conn:
        assert conn.execute(text('SELECT order_date FROM sales_order')).scalar_one() == date(2024, 6, 15)
    # Full-form save must also write only the current source line's date.
    from app.routers.orders import OrderUpdate
    with db() as conn:
        from app.financial_calculations import line_snapshot
        source = line_snapshot(conn, ids['设备-2025'])
    aliases = {'amount_type':'gross_net_type', 'statistical_category':'statistic_category',
               'team_name':'team_level3_name', 'user_name':'end_user_name',
               'net_unit_price':'sales_unit_price_no_tax', 'unit_price':'sales_unit_price',
               'net_revenue':'revenue_no_tax'}
    payload = {key: source.get(aliases.get(key, key)) for key in OrderUpdate.model_fields}
    payload['order_date'] = '2026-01-02'
    import json
    response = client.put(f"/api/orders/{ids['设备-2025']}", json=json.loads(json.dumps(payload, default=str)), headers=headers)
    assert response.status_code == 200, response.text
    rows = client.get('/api/orders', headers=headers).json()['items']
    assert {r['goods_name']: r['order_date'] for r in rows} == {
        '设备-2025': '2026-01-02', '设备-2024': '2024-09-18',
    }


def test_regular_orders_still_use_shared_order_date(client, headers):
    from test_financial_integration import _create_order

    line_id, _ = _create_order(client, headers, 'REGULAR-DATE')
    item = batch_item(client, headers, 'orders', line_id)
    item['order_date'] = '2024-12-31'
    response = client.put('/api/orders/batch-basic', json={'items': [item]}, headers=headers)
    assert response.status_code == 200, response.text
    rows = client.get('/api/orders', params={'start_date':'2024-12-31', 'end_date':'2024-12-31'}, headers=headers).json()['items']
    assert len(rows) == 1 and rows[0]['order_date'] == '2024-12-31'


def test_restore_backup_from_before_line_dates_recovers_original_dates(client, headers):
    import gzip
    import json
    from test_backup_integrity import backup_file, rewrite

    import_dated_line(client, headers, 2025, 200)
    import_dated_line(client, headers, 2024, 100)
    backup_id, path = backup_file(client, headers)
    with gzip.open(path, 'rt', encoding='utf-8') as stream:
        payload = json.load(stream)
    for row in payload['tables']['order_line']:
        row.pop('line_order_date')
        row.pop('line_order_date_initialized')
    rewrite(path, payload)
    response = client.post(f'/api/backups/{backup_id}/restore', headers=headers)
    assert response.status_code == 200, response.text
    rows = client.get('/api/orders', headers=headers).json()['items']
    assert {r['goods_name']:r['order_date'] for r in rows} == {
        '设备-2025':'2025-06-15', '设备-2024':'2024-06-15',
    }
